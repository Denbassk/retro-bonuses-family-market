import openpyxl, csv
from collections import defaultdict

wb = openpyxl.load_workbook(r"D:\РЕТРО_БОНУСЫ Фэмэли маркет\Ассортиментная матрица.xlsx", data_only=True)
ws = wb.active
agg = defaultdict(lambda: {'cnt': 0, 'sample': ''})
header = [c.value for c in ws[1]]
i_sup = header.index('Поставщик')
i_prod = header.index('Товар')
for row in ws.iter_rows(min_row=2, values_only=True):
    sup = row[i_sup]
    if not sup or 'ИТОГО' in str(sup):
        continue
    agg[sup]['cnt'] += 1
    if not agg[sup]['sample']:
        agg[sup]['sample'] = row[i_prod]

with open(r"D:\РЕТРО_БОНУСЫ Фэмэли маркет\matrix_mapping_template.csv", 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['matrix_supplier','sku_count','sample_product','db_supplier','db_brand','action'])
    for sup, d in sorted(agg.items(), key=lambda x: -x[1]['cnt']):
        w.writerow([sup, d['cnt'], d['sample'], '', '', ''])

print(f"Saved {len(agg)} matrix suppliers")
