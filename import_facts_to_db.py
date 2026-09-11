#!/usr/bin/env python3
"""
import_facts_to_db.py - ЭТАП 3: запись фактов оплат ретро в Supabase.

Источники: лист «2026» файла «Ретро Бонусы.xlsx» + примечания к ячейкам + retro_fact_name_map.
По умолчанию - пробный прогон. Запись только с --apply.

Правила безопасности:
  - is_ignored и составные (split_targets) НЕ пишутся автоматически;
  - при изменении ранее внесённой суммы прежнее значение уходит в amount_previous;
  - подозрительные даты не записываются, ячейка помечается needs_manual.

Запуск:
  python import_facts_to_db.py "Ретро Бонусы.xlsx" --sheet 2026
  python import_facts_to_db.py "Ретро Бонусы.xlsx" --sheet 2026 --apply
"""
import re, argparse, getpass
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

import sb
from import_retro_facts import norm_name, parse_amount, classify_header, STOP_NAMES
from parse_payment_notes import parse_note, AUTHOR_RE, DATE_RE, EARMARK_RE
from parse_note_components import components, month_refs, ADJ_WORDS
from dump_cell_comments import comments_from_zip

NON_RETRO = ("дмп", "кубы", "кубі", "куби", "стойки", "стойка", "сгущёнка", "сгущенка")
SRC = "excel_retro_sheet"


def note_body(text):
    b = AUTHOR_RE.sub("", re.sub(r"\s+", " ", text or "").strip())
    b = EARMARK_RE.sub(" ", b)
    return DATE_RE.sub(" ", b).strip(" .,-—")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    path = Path(args.xlsx)
    year = int(re.search(r"(20\d{2})", args.sheet).group(1))
    ws = load_workbook(path, data_only=True)[args.sheet]
    who = getpass.getuser()
    now = datetime.now(timezone.utc).isoformat()

    cols = {}
    for c in range(2, ws.max_column + 1):
        cl = classify_header(ws.cell(row=args.header_row, column=c).value, year)
        if cl and cl[0] != "skip":
            cols[c] = cl

    cmts = {}
    for r in range(args.header_row + 1, ws.max_row + 1):
        for c in cols:
            cm = ws.cell(row=r, column=c).comment
            if cm and (cm.text or "").strip():
                cmts[(r, c)] = cm.text
    if not cmts:
        for ref, txt in comments_from_zip(path, args.sheet).items():
            m = re.match(r"([A-Z]+)(\d+)", ref)
            if m:
                cmts[(int(m.group(2)), column_index_from_string(m.group(1)))] = txt

    nmap = {m["excel_name_normalized"]: m for m in sb.get(
        "retro_fact_name_map",
        "select=excel_name_normalized,supplier_id,is_ignored,split_targets,needs_manual")}
    sup_name = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}
    existing = {(x["supplier_id"], x["period_label"]): x for x in sb.get(
        "retro_payments_fact", "select=supplier_id,period_label,amount_paid,payment_date")}

    facts, openings, skipped, manual, unmapped = [], [], [], [], set()

    for r in range(args.header_row + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=1).value
        if raw is None or not str(raw).strip():
            continue
        raw = str(raw).strip()
        nm = norm_name(raw)
        if nm in STOP_NAMES:
            continue

        mp = nmap.get(nm)
        if not mp:
            unmapped.add(raw)
            continue
        if mp.get("is_ignored"):
            skipped.append(raw)
            continue
        if mp.get("split_targets"):
            manual.append((raw, "составной платёж (split_targets)"))
            continue
        sid = mp.get("supplier_id")
        if not sid:
            unmapped.add(raw)
            continue

        for c, (kind, label) in cols.items():
            amt = parse_amount(ws.cell(row=r, column=c).value)
            if amt is None:
                continue
            cell = f"{get_column_letter(c)}{r}"
            txt = cmts.get((r, c), "")
            p = parse_note(txt, label if kind == "month" else None, fallback_year=year)
            body = note_body(txt)

            if kind == "one_off":
                openings.append({
                    "supplier_id": sid, "store_label": label, "amount": amt, "year": year,
                    "payment_date": p["payment_date"].isoformat() if p["payment_date"] else None,
                    "source_file": path.name, "imported_at": now})
                continue

            flags = []
            if p["date_warning"]:
                flags.append(p["date_warning"])
            if mp.get("needs_manual"):
                flags.append("поставщик помечен как ручной")

            extra = [{"amount": e["amount"], "label": e["label"], "source": "note"}
                     for e in p["earmarks"]]
            covers = []
            if body:
                comps = [x for x in components(body) if x["is_money"]]
                for x in comps:
                    if any(x["label"].startswith(k) for k in NON_RETRO):
                        extra.append({"amount": x["amount"], "label": x["label"],
                                      "source": "note"})
                own = int(label.split("-")[1])
                other = sorted(month_refs(body) - {own})
                if other and any(w in body.lower() for w in ADJ_WORDS):
                    covers = [f"{year}-{m:02d}" for m in other]
                s = sum(x["amount"] for x in comps)
                if comps and len(comps) >= 2 and abs(s - amt) > 5:
                    flags.append(f"состав примечания {s:,.0f} != суммы {amt:,.0f}")

            prev = existing.get((sid, label))
            if prev and abs(float(prev["amount_paid"]) - amt) < 0.005:
                continue

            rec = {"supplier_id": sid, "period_label": label, "amount_paid": amt,
                   "payment_date": (p["payment_date"].isoformat()
                                    if p["payment_date"] and not p["date_warning"] else None),
                   "notes": txt or None, "source_file": path.name, "import_source": SRC,
                   "imported_at": now, "needs_manual": bool(flags),
                   "covers_periods": covers or None,
                   "additional_payments": extra}
            if prev:
                rec["amount_previous"] = float(prev["amount_paid"])
            facts.append((cell, raw, rec, flags))

    ins = [x for x in facts if "amount_previous" not in x[2]]
    upd = [x for x in facts if "amount_previous" in x[2]]
    print(f"[i] Новых записей: {len(ins)} | обновлений: {len(upd)} | "
          f"разовых бонусов: {len(openings)}")
    print(f"    пропущено (ignore): {len(skipped)} | составных: {len(manual)} | "
          f"без маппинга: {len(unmapped)}")

    if unmapped:
        print("\n[!] Нет в retro_fact_name_map:")
        for x in sorted(unmapped):
            print(f"    {x}")
    if manual:
        print("\n[~] Требуют ручного разбора:")
        for x, why in manual:
            print(f"    {x} — {why}")

    if upd:
        print("\n[!] Изменение ранее внесённых сумм:")
        for cell, raw, rec, _ in upd:
            d = rec["amount_paid"] - rec["amount_previous"]
            print(f"    {cell:<6} {raw[:26]:<26} {rec['period_label']}  "
                  f"{rec['amount_previous']:,.0f} -> {rec['amount_paid']:,.0f} ({d:+,.0f})")

    flagged = [x for x in facts if x[3]]
    if flagged:
        print(f"\n[~] Помечены needs_manual ({len(flagged)}):")
        for cell, raw, rec, fl in flagged:
            print(f"    {cell:<6} {raw[:26]:<26} {rec['period_label']}  {'; '.join(fl)}")

    withextra = [x for x in facts if x[2]["additional_payments"]]
    if withextra:
        print(f"\n[i] Выделены нефакторные компоненты ({len(withextra)}):")
        for cell, raw, rec, _ in withextra:
            e = ", ".join(f"{a['amount']:,.0f} {a['label']}" for a in rec["additional_payments"])
            print(f"    {cell:<6} {raw[:26]:<26} {rec['period_label']}  {e}")

    withcov = [x for x in facts if x[2]["covers_periods"]]
    if withcov:
        print(f"\n[i] Платежи за несколько периодов ({len(withcov)}):")
        for cell, raw, rec, _ in withcov:
            print(f"    {cell:<6} {raw[:26]:<26} {rec['period_label']} "
                  f"+ {', '.join(rec['covers_periods'])}")

    if not args.apply:
        print("\n[i] Пробный прогон. Для записи добавьте --apply")
        return

    if facts:
        sb.upsert("retro_payments_fact", [x[2] for x in facts],
                  on_conflict="supplier_id,period_label")
        print(f"\n[>] retro_payments_fact: записано {len(facts)}")
    if openings:
        sb.upsert("store_opening_bonuses", openings,
                  on_conflict="supplier_id,store_label,year")
        print(f"[>] store_opening_bonuses: записано {len(openings)}")


if __name__ == "__main__":
    main()
