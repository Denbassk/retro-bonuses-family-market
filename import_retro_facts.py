#!/usr/bin/env python3
"""
import_retro_facts.py - ЭТАП 1: чтение факт-таблицы ретро из Excel.
Ничего не пишет в Supabase. Только парсит и выгружает CSV для проверки.

Запуск:
  python import_retro_facts.py "Ретро Бонусы.xlsx" --sheet 2026
"""
import re, csv, argparse
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook

MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5, "июнь": 6,
    "июль": 7, "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
    "січень": 1, "лютий": 2, "березень": 3, "квітень": 4, "травень": 5, "червень": 6,
    "липень": 7, "серпень": 8, "вересень": 9, "жовтень": 10, "листопад": 11, "грудень": 12,
}
STOP_NAMES = {"общая сумма", "сумма", "итого", "всего", "разом", "total"}
UA_RU = str.maketrans({"і": "и", "ї": "и", "є": "е", "ґ": "г", "’": "'", "`": "'"})


def norm_name(s):
    """Каноническое имя для сопоставления: регистр, укр/рус буквы, пробелы, скобки."""
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ").lower().translate(UA_RU)
    s = re.sub(r"[«»\"']", "", s)
    s = re.sub(r"\s*\(\s*", " (", s)
    s = re.sub(r"\s*\)", ")", s)
    return re.sub(r"\s+", " ", s).strip()


def parse_amount(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def classify_header(h, year):
    """month -> период YYYY-MM, one_off -> разовый бонус, skip -> служебная колонка."""
    if h is None:
        return None
    t = re.sub(r"\s+", " ", str(h).replace("\xa0", " ").strip().lower())
    if t in ("поставщик", "условия", ""):
        return ("skip", t)
    key = t.rstrip(".")
    if key in MONTHS:
        return ("month", f"{year}-{MONTHS[key]:02d}")
    return ("one_off", t)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    path = Path(args.xlsx)
    year = int(re.search(r"(20\d{2})", args.sheet).group(1))
    ws = load_workbook(path, data_only=True)[args.sheet]

    cols = {}
    for c in range(2, ws.max_column + 1):
        cl = classify_header(ws.cell(row=args.header_row, column=c).value, year)
        if cl and cl[0] != "skip":
            cols[c] = cl
    months = [v[1] for v in cols.values() if v[0] == "month"]
    oneoffs = [v[1] for v in cols.values() if v[0] == "one_off"]
    print(f"[i] Месяцев: {len(months)} -> {', '.join(sorted(months))}")
    print(f"[i] Разовых колонок: {len(oneoffs)} -> {', '.join(oneoffs) or '-'}")

    facts, totals_row, seen = [], {}, defaultdict(list)
    for r in range(args.header_row + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=1).value
        if raw is None or not str(raw).strip():
            continue
        raw = str(raw).strip()
        nm = norm_name(raw)

        if nm in STOP_NAMES:
            for c, (kind, label) in cols.items():
                a = parse_amount(ws.cell(row=r, column=c).value)
                if a is not None:
                    totals_row[label] = a
            continue

        seen[nm].append(r)
        cond = ws.cell(row=r, column=2).value
        for c, (kind, label) in cols.items():
            amount = parse_amount(ws.cell(row=r, column=c).value)
            facts.append({
                "row": r, "raw_name": raw, "norm_name": nm,
                "condition_raw": "" if cond is None else str(cond),
                "kind": kind,
                "period_label": label if kind == "month" else "",
                "oneoff_label": label if kind == "one_off" else "",
                "amount": "" if amount is None else f"{amount:.2f}",
                "state": "empty" if amount is None else ("zero" if amount == 0 else "value"),
            })

    dups = {k: v for k, v in seen.items() if len(v) > 1}
    if dups:
        print(f"\n[!] ДУБЛИ имён ({len(dups)}):")
        for k, rows in dups.items():
            print(f"    {k}  -> строки {rows}")

    if totals_row:
        print("\n[i] Сверка с итоговой строкой:")
        by_label = defaultdict(float)
        for f in facts:
            if f["amount"]:
                by_label[f["period_label"] or f["oneoff_label"]] += float(f["amount"])
        for label, total in sorted(totals_row.items()):
            calc = by_label.get(label, 0.0)
            d = calc - total
            mark = "OK " if abs(d) < 1 else "!! "
            print(f"  {mark}{label:<16} итог={total:>14,.2f}  сумма строк={calc:>14,.2f}  дельта={d:>12,.2f}")
    else:
        print("\n[i] Итоговая строка не найдена.")

    vals = sum(1 for f in facts if f["state"] == "value")
    zeros = sum(1 for f in facts if f["state"] == "zero")
    empty = sum(1 for f in facts if f["state"] == "empty")
    print(f"\n[i] Поставщиков: {len(seen)} | ячеек: значений {vals}, нулей {zeros}, пустых {empty}")

    out = Path(args.out or f"retro_facts_{args.sheet}.csv")
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(facts[0].keys()))
        w.writeheader()
        w.writerows(facts)
    print(f"[>] {out}  ({len(facts)} строк)")


if __name__ == "__main__":
    main()
