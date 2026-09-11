#!/usr/bin/env python3
"""
load_name_map.py - заливка ручного маппинга имён в retro_fact_name_map.

Источник: "Маппинг имен.xlsx" (колонки: имя в таблице ретро | имя в HTML-панели | примечание).
"Не активный" -> is_ignored. Несколько имён через ";" -> split_targets (ручной разбор).

Запуск:
  python load_name_map.py "Маппинг имен.xlsx" --facts "Ретро Бонусы.xlsx" --sheet 2026
"""
import re, argparse, getpass
from pathlib import Path
from openpyxl import load_workbook

import sb
from import_retro_facts import norm_name, STOP_NAMES
from resolve_supplier_names import score

INACTIVE = {"не активный", "не активний", "неактивный", "-", ""}
NOT_SUPPLIER = {"оплата безнал", "безнал"}
MANUAL_HINT = ("не будут сходит", "правки руками", "компенсац")


def find_col(ws, *keys):
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(row=1, column=c).value or "").lower()
        if any(k in h for k in keys):
            return c
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mapping")
    ap.add_argument("--facts", default=None, help="файл Ретро Бонусы.xlsx для кросс-проверки")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--apply", action="store_true", help="записать в БД")
    args = ap.parse_args()

    ws = load_workbook(Path(args.mapping), data_only=True).worksheets[0]
    c_ex = find_col(ws, "таблице ретро", "таблиці ретро")
    c_ad = find_col(ws, "html", "панел")
    c_nt = find_col(ws, "римичан", "римечан", "имітк")
    if not c_ex or not c_ad:
        raise SystemExit(f"[!] Не нашёл колонки. excel={c_ex}, admin={c_ad}")

    suppliers = sb.get("suppliers", "select=id,name")
    sup_norm = {s["id"]: norm_name(s["name"]) for s in suppliers}
    by_norm = {}
    for s in suppliers:
        by_norm.setdefault(sup_norm[s["id"]], s)

    who = getpass.getuser()
    recs, problems, splits = [], [], []

    for r in range(2, ws.max_row + 1):
        ex_raw = ws.cell(row=r, column=c_ex).value
        if ex_raw is None or not str(ex_raw).strip():
            continue
        ex_raw = str(ex_raw).strip()
        nm = norm_name(ex_raw)
        if nm in STOP_NAMES:
            continue
        ad_raw = str(ws.cell(row=r, column=c_ad).value or "").strip()
        note = str(ws.cell(row=r, column=c_nt).value or "").strip() if c_nt else ""
        note = "" if note.lower() == "nan" else note
        manual = any(h in note.lower() for h in MANUAL_HINT)

        rec = {"excel_name_normalized": nm, "excel_name_raw": ex_raw,
               "admin_name_raw": ad_raw, "notes": note or None,
               "needs_manual": manual, "is_ignored": False,
               "supplier_id": None, "split_targets": None, "confirmed_by": who}

        low = ad_raw.lower().strip()
        if low in INACTIVE or low in NOT_SUPPLIER:
            rec["is_ignored"] = True
            rec["notes"] = (note + " | " if note else "") + f"не сверяется: {ad_raw}"
            recs.append(rec)
            continue

        parts = [p.strip() for p in ad_raw.split(";") if p.strip()]
        resolved = []
        for p in parts:
            s = by_norm.get(norm_name(p))
            if s:
                resolved.append({"supplier_id": s["id"], "supplier_name": s["name"]})
            else:
                best = sorted(((score(norm_name(p), sup_norm[x["id"]]), x)
                               for x in suppliers), key=lambda z: -z[0])[:3]
                problems.append((ex_raw, p, best))
                resolved.append(None)

        if any(x is None for x in resolved):
            continue
        if len(resolved) == 1:
            rec["supplier_id"] = resolved[0]["supplier_id"]
        else:
            rec["split_targets"] = resolved
            rec["needs_manual"] = True
            splits.append((ex_raw, [x["supplier_name"] for x in resolved]))
        recs.append(rec)

    print(f"[i] Готово к записи: {len(recs)}")
    print(f"    из них ignore: {sum(1 for x in recs if x['is_ignored'])}, "
          f"ручных: {sum(1 for x in recs if x['needs_manual'])}, "
          f"составных: {len(splits)}")
    for ex, names in splits:
        print(f"    SPLIT: {ex} -> {', '.join(names)}")

    if problems:
        print(f"\n[!] Не найдены в справочнике ({len(problems)}):")
        for ex, p, best in problems:
            print(f"    «{p}»  (из строки «{ex}»)")
            for scv, s in best:
                print(f"        похоже [{scv:.2f}] {s['name']}")

    if args.facts:
        fws = load_workbook(Path(args.facts), data_only=True)[args.sheet]
        sheet_names = set()
        for r in range(2, fws.max_row + 1):
            v = fws.cell(row=r, column=1).value
            if v and str(v).strip():
                n = norm_name(str(v).strip())
                if n not in STOP_NAMES:
                    sheet_names.add(n)
        mapped = {x["excel_name_normalized"] for x in recs}
        miss = sheet_names - mapped
        extra = mapped - sheet_names
        print(f"\n[i] Кросс-проверка с листом «{args.sheet}»:")
        print(f"    имён в листе: {len(sheet_names)}, покрыто маппингом: {len(sheet_names & mapped)}")
        for n in sorted(miss):
            print(f"    [!] в листе есть, в маппинге нет: {n}")
        for n in sorted(extra):
            print(f"    [~] в маппинге есть, в листе нет: {n}")

    if not args.apply:
        print("\n[i] Пробный прогон. Для записи добавьте --apply")
        return
    if problems:
        print("\n[!] Есть нерешённые имена — запись отменена.")
        return
    sb.upsert("retro_fact_name_map", recs, on_conflict="excel_name_normalized")
    print(f"\n[>] Записано: {len(recs)}")


if __name__ == "__main__":
    main()
