#!/usr/bin/env python3
"""
migrate_adjustment_modes.py - разовая миграция доплат «не ретро» после sql/2026-09-11_adjustments_payment_mode.sql.

1. Доп. выплаты, записанные старой админкой в retro_payments_fact.additional_payments (Глобал-Сервіс 02-04,
   Боржомі 03), -> retro_adjustments in_payment (деньги в «Оплачено», по числам факт - расчёт = доплата).
   Пропуск: сумма пустая, <= 0 или больше самого факта (Глобал-Сервіс 01: 80 245 при факте 15 172 - мусор).
2. Юрія 2026-07 «сгущёнка» 1 289 - это РЕТРО по правилу 15%, из additional_payments убирается.
3. ДМП БІР Славутич: режим решают числа Excel. Если ячейка группы (Кег + Славутич) = факты админки + ДМП,
   значит ДМП в «Оплачено» не внесён -> separate (заплачено отдельно). Май: существующая доплата переводится,
   июнь: добавляется separate, если в примечании ячейки есть ДМП и разница Excel - админка = его сумме.

Запуск (из корня):  python tools\\migrate_adjustment_modes.py          # показать
                    python tools\\migrate_adjustment_modes.py --apply
"""
import re, sys, argparse
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "core"))
import sb
from adjustments import add_adjustment, set_mode
from reconcile_facts import newest_excel, load_excel
from dump_cell_comments import comments_from_zip
from import_facts_to_db import note_body
from parse_note_components import components
from openpyxl import load_workbook

RETRO_WORDS = ("сгущ",)
SLAV, KEG = "БІР (Пиво Славутич)", "БІР (Пиво Кег)"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    try:
        sb.get("retro_adjustments", "select=payment_mode&limit=1")
    except Exception as e:
        if a.apply:
            sys.exit(f"[!] нет колонки payment_mode - сначала выполните sql/2026-09-11_adjustments_payment_mode.sql ({e})")
        print("[!] колонки payment_mode ещё нет (нужен SQL) - показываю, что будет сделано\n")

    sup = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}
    name2id = {v: k for k, v in sup.items()}
    adjs = sb.get("retro_adjustments", "select=*")

    print("[1] доп. выплаты из фактов -> retro_adjustments (in_payment)")
    for f in sb.get("retro_payments_fact", "select=id,supplier_id,period_label,amount_paid,additional_payments"):
        ap_list = f.get("additional_payments") or []
        if not ap_list:
            continue
        keep, name = [], sup.get(f["supplier_id"], "?")
        for e in ap_list:
            amt, label = e.get("amount"), (e.get("label") or e.get("note") or "доп. выплата").strip()
            paid = float(f["amount_paid"] or 0)
            if any(w in label.lower() for w in RETRO_WORDS):
                print(f"    {name} {f['period_label']}: «{label}» {amt} - это ретро по правилу, убираю из доп. выплат")
                continue
            if not amt or float(amt) <= 0 or float(amt) > paid:
                print(f"    {name} {f['period_label']}: «{label}» {amt} при факте {paid:,.0f} - пропуск (нет суммы или больше факта)")
                keep.append(e)
                continue
            dup = [x for x in adjs if x["supplier_id"] == f["supplier_id"] and x["period_label"] == f["period_label"]
                   and abs(float(x["amount"]) - float(amt)) < 0.005]
            if dup:
                print(f"    {name} {f['period_label']}: {float(amt):,.2f} уже в retro_adjustments")
                continue
            res = add_adjustment(f["supplier_id"], f["period_label"], float(amt), f"{label} (перенесено из доп. выплаты факта)",
                                 source="migration_fact_extra", created_by="migrate_adjustment_modes.py", apply=a.apply)
            print(f"    {name} {f['period_label']}: +{float(amt):,.2f} «{label}» -> {res}")
        if len(keep) != len(ap_list) and a.apply:
            sb._req(f"{sb.URL}/rest/v1/retro_payments_fact?id=eq.{f['id']}", {"additional_payments": keep},
                    "PATCH", {"Prefer": "return=minimal"})

    print("\n[2] ДМП БІР Славутич: Excel группы против фактов админки")
    slav, keg = name2id.get(SLAV), name2id.get(KEG)
    if not slav or not keg:
        print(f"    нет поставщиков {SLAV} / {KEG} - пропуск")
        return
    nmap = {m["excel_name_normalized"]: m for m in sb.get(
        "retro_fact_name_map", "select=excel_name_normalized,supplier_id,is_ignored,split_targets,needs_manual")}
    xlsx = newest_excel()
    excel = load_excel(xlsx, "2026", nmap, name2id)
    wb = load_workbook(xlsx, data_only=True)["2026"]
    cm = {f"{c.column_letter}{c.row}": c.comment.text for row in wb.iter_rows() for c in row if c.comment}
    cm = cm or comments_from_zip(xlsx, "2026")
    facts = {(f["supplier_id"], f["period_label"]): float(f["amount_paid"] or 0)
             for f in sb.get("retro_payments_fact", f"supplier_id=in.({slav},{keg})&select=supplier_id,period_label,amount_paid")}
    key = tuple(sorted((slav, keg)))
    for per in ("2026-05", "2026-06", "2026-07", "2026-08"):
        x, cell = excel.get((key, per), (None, ""))
        adm = facts.get((slav, per), 0.0) + facts.get((keg, per), 0.0)
        note = note_body(cm.get(cell, ""))
        dmp = [c for c in components(note) if c["is_money"] and c["label"].lower().startswith("дмп")] if note else []
        have = [j for j in adjs if j["supplier_id"] == slav and j["period_label"] == per and "дмп" in (j["notes"] or "").lower()]
        gap = (x - adm) if x is not None else None
        print(f"    {per}: Excel {cell} = {x}, админка Кег+Славутич = {adm:,.0f}, разница {gap if gap is None else round(gap, 2)}"
              f" | примечание: «{note[:60]}» | ДМП в доплатах: {[float(j['amount']) for j in have]}")
        if gap is None:
            continue
        for j in have:
            amt = float(j["amount"])
            if abs(gap - amt) < 1 and (j.get("payment_mode") or "in_payment") != "separate":
                print(f"      -> ДМП {amt:,.0f} не входит в «Оплачено» админки: separate: {set_mode(j['id'], 'separate', apply=a.apply)}")
            elif abs(gap) < 1:
                print(f"      -> Excel = админка, ДМП {amt:,.0f} уже в «Оплачено»: оставляю in_payment")
        if not have and dmp:
            amt = dmp[0]["amount"]
            if abs(gap - amt) < 1:
                print(f"      -> добавляю ДМП {amt:,.0f} separate: "
                      f"{add_adjustment(slav, per, amt, 'ДМП', source='migration_excel_note', created_by='migrate_adjustment_modes.py', payment_mode='separate', apply=a.apply)}")
            else:
                print(f"      -> ДМП {amt:,.0f} в примечании, но разница Excel - админка {gap:,.0f} - решать вручную")
    if not a.apply:
        print("\n[i] Пробный прогон. Для записи добавьте --apply")


if __name__ == "__main__":
    main()
