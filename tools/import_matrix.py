"""
Импорт ассортиментной матрицы 2026 в Supabase.brand_barcodes
"""
import os, sys, re
from pathlib import Path
from openpyxl import load_workbook
from supabase import create_client
from dotenv import load_dotenv

load_dotenv()
SUPABASE_URL = os.getenv('SUPABASE_URL')
SUPABASE_KEY = os.getenv('SUPABASE_SERVICE_KEY')
sb = create_client(SUPABASE_URL, SUPABASE_KEY)

XLSX_PATH = Path(r"D:\РЕТРО_БОНУСЫ Фэмэли маркет\data\справочники\Ассортиментная матрица.xlsx")

# ---------- нормализация ----------
def normalize(s):
    if not s:
        return ''
    s = str(s).lower().strip()
    table = str.maketrans({'і':'и','ї':'и','є':'е','ё':'е','ы':'и','ъ':'','ь':'',"'":'','`':'','*':''})
    s = s.translate(table)
    s = re.sub(r'[^a-zа-я0-9 ]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s

# ---------- чтение xlsx ----------
def read_matrix():
    wb = load_workbook(XLSX_PATH, data_only=True)
    ws = wb.active
    rows = []
    headers = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(h).strip() if h else '' for h in row]
            continue
        if not row or not any(row):
            continue
        rec = dict(zip(headers, row))
        sup = (rec.get('Поставщик') or '').strip()
        name = (rec.get('Товар') or '').strip()
        cat = (rec.get('Категория') or '').strip()
        decision = (rec.get('Решение') or '').strip()
        bc = rec.get('Штрихкод')
        # игнорим строки ИТОГО
        if not sup or not name or not bc:
            continue
        if 'ИТОГО' in name.upper():
            continue
        # штрихкод как строка без точек/запятых
        bc_str = str(bc).strip()
        bc_str = bc_str.replace(',', '').replace('.', '').replace(' ', '')
        if not bc_str.isdigit():
            continue
        rows.append({
            'matrix_supplier': sup,
            'product_name': name,
            'category': cat,
            'decision': decision,
            'barcode': bc_str,
        })
    return rows

# ---------- загрузка справочников ----------
def load_refs():
    suppliers = sb.table('suppliers').select('id,name').execute().data
    brands = sb.table('supplier_brands').select('id,name,supplier_id').execute().data
    aliases = sb.table('supplier_aliases').select('alias_name,supplier_id,supplier_brand_id').execute().data

    sup_by_norm = {normalize(s['name']): s for s in suppliers}
    brands_by_sup = {}
    for b in brands:
        brands_by_sup.setdefault(b['supplier_id'], []).append(b)
    brand_by_norm = {}
    for b in brands:
        brand_by_norm.setdefault(normalize(b['name']), []).append(b)
    alias_by_norm = {}
    for a in aliases:
        if a.get('supplier_id') or a.get('supplier_brand_id'):
            alias_by_norm[normalize(a['alias_name'])] = a
    return sup_by_norm, brands_by_sup, brand_by_norm, alias_by_norm

def resolve_matrix_row(matrix_supplier, product_name,
                       sup_by_norm, brands_by_sup, brand_by_norm, alias_by_norm):
    """
    Возвращает (supplier_brand_id, matched_by, debug_reason)
    """
    n = normalize(matrix_supplier)
    pn_norm = ' ' + normalize(product_name) + ' '

    # 1. Прямое совпадение матричного "поставщика" с именем бренда в БД
    if n in brand_by_norm:
        cands = brand_by_norm[n]
        if len(cands) == 1:
            return cands[0]['id'], 'brand_name_exact', None
        # несколько брендов с тем же именем у разных supplier-ов — берём первого
        return cands[0]['id'], 'brand_name_exact_multi', None

    # 2. Алиас → если ведёт на supplier_brand_id напрямую
    if n in alias_by_norm:
        a = alias_by_norm[n]
        if a.get('supplier_brand_id'):
            return a['supplier_brand_id'], 'alias_brand', None
        # ведёт на supplier — пойдём искать бренд по product_name
        sid = a['supplier_id']
        bid, how = match_brand_in_pn(pn_norm, brands_by_sup.get(sid, []))
        if bid:
            return bid, f'alias_supplier_{how}', None

    # 3. Совпадение с supplier-ом
    if n in sup_by_norm:
        sid = sup_by_norm[n]['id']
        brands = brands_by_sup.get(sid, [])
        bid, how = match_brand_in_pn(pn_norm, brands)
        if bid:
            return bid, f'supplier_{how}', None
        return None, None, f"supplier '{matrix_supplier}' OK, у него {len(brands)} брендов: {', '.join(b['name'] for b in brands)[:150]} — keyword не найден"

    # 4. Поиск имени бренда внутри product_name (на случай когда матричное имя нерелевантно)
    candidates = []
    for bn, blist in brand_by_norm.items():
        if not bn or len(bn) < 4:
            continue
        if ' ' + bn + ' ' in pn_norm or pn_norm.startswith(bn + ' '):
            for b in blist:
                candidates.append((len(bn), b))
    if candidates:
        candidates.sort(key=lambda x: -x[0])
        return candidates[0][1]['id'], 'pn_brand_search', None

    return None, None, f"matrix_supplier '{matrix_supplier}' не найден ни в suppliers, ни в supplier_brands, ни в aliases"


def match_brand_in_pn(pn_norm, brands):
    if not brands:
        return None, None
    if len(brands) == 1:
        return brands[0]['id'], 'single_brand'
    matches = []
    for b in brands:
        bn = normalize(b['name'])
        if not bn or len(bn) < 3:
            continue
        if ' ' + bn + ' ' in pn_norm or pn_norm.startswith(bn + ' ') or pn_norm.endswith(' ' + bn):
            matches.append((len(bn), b))
    if matches:
        matches.sort(key=lambda x: -x[0])
        return matches[0][1]['id'], 'keyword'
    return None, None

def match_brand(product_name, brands):
    """ищем бренд по вхождению нормализованного имени в product_name"""
    if not brands:
        return None, None
    if len(brands) == 1:
        return brands[0]['id'], 'single_brand'
    pn = ' ' + normalize(product_name) + ' '
    matches = []
    for b in brands:
        bn = normalize(b['name'])
        if not bn or len(bn) < 3:
            continue
        if ' ' + bn + ' ' in pn or pn.startswith(bn + ' ') or pn.endswith(' ' + bn):
            matches.append((len(bn), b))
    if matches:
        matches.sort(key=lambda x: -x[0])
        return matches[0][1]['id'], 'keyword'
    return None, None

# ---------- main ----------
def main():
    print(f"Чтение {XLSX_PATH.name}...")
    rows = read_matrix()
    print(f"Прочитано строк (без ИТОГО): {len(rows)}")

    sup_by_norm, sup_by_id, brands_by_sup, alias_by_norm = load_refs()
    print(f"В БД: {len(sup_by_norm)} поставщиков, {sum(len(v) for v in brands_by_sup.values())} брендов")

    barcodes_to_load = {}     # barcode -> dict
    unmapped = {}             # barcode -> dict
    stats = {'single_brand': 0, 'keyword': 0, 'no_supplier': 0, 'no_brand': 0, 'duplicate': 0}
    by_unmapped_sup = {}

    sup_by_norm, brands_by_sup, brand_by_norm, alias_by_norm = load_refs()
    print(f"В БД: {len(sup_by_norm)} поставщиков, "
          f"{sum(len(v) for v in brands_by_sup.values())} брендов, "
          f"{sum(len(v) for v in brand_by_norm.values())} brand-name записей")

    barcodes_to_load = {}
    unmapped = {}
    stats = {}
    by_unmapped_sup = {}

    for r in rows:
        bc = r['barcode']
        if bc in barcodes_to_load or bc in unmapped:
            stats['duplicate'] = stats.get('duplicate', 0) + 1
            unmapped.setdefault(bc, {
                'matrix_supplier': r['matrix_supplier'],
                'barcode': bc, 'product_name': r['product_name'],
                'category': r['category'], 'reason': 'duplicate barcode',
            })
            continue

        bid, how, reason = resolve_matrix_row(
            r['matrix_supplier'], r['product_name'],
            sup_by_norm, brands_by_sup, brand_by_norm, alias_by_norm
        )
        if bid:
            stats[how] = stats.get(how, 0) + 1
            barcodes_to_load[bc] = {
                'supplier_brand_id': bid,
                'barcode': bc,
                'product_name': r['product_name'],
                'category': r['category'],
                'is_active': True,
                'source': 'matrix_2026',
                'matched_by': how,
            }
        else:
            stats['unmapped'] = stats.get('unmapped', 0) + 1
            by_unmapped_sup[r['matrix_supplier']] = by_unmapped_sup.get(r['matrix_supplier'], 0) + 1
            unmapped[bc] = {
                'matrix_supplier': r['matrix_supplier'],
                'barcode': bc, 'product_name': r['product_name'],
                'category': r['category'], 'reason': reason or 'unknown',
            }

    print(f"\n=== Итоги маппинга ===")
    for k, v in sorted(stats.items(), key=lambda x: -x[1]):
        print(f"  {k:30s} {v}")
    print(f"  → к загрузке brand_barcodes:    {len(barcodes_to_load)}")
    print(f"  → к загрузке unmapped:          {len(unmapped)}")

    print(f"\n=== Топ-15 матричных поставщиков с проблемами ===")
    top = sorted(by_unmapped_sup.items(), key=lambda x: -x[1])[:15]
    for sup, cnt in top:
        print(f"  {sup}: {cnt} SKU")

    print(f"\n--- Превью замапленных (первые 10) ---")
    for v in list(barcodes_to_load.values())[:10]:
        print(f"  {v['barcode']} → brand={v['supplier_brand_id'][:8]}... ({v['matched_by']}) — {v['product_name'][:60]}")

    print(f"\n--- Превью unmapped (первые 10) ---")
    for v in list(unmapped.values())[:10]:
        print(f"  {v['barcode']} [{v['matrix_supplier']}] {v['product_name'][:50]} — {v['reason'][:80]}")

    ans = input("\nЗагрузить в Supabase? (y/n): ").strip().lower()
    if ans != 'y':
        print("Отменено.")
        return

    # batch upsert (по 500)
    def chunks(lst, n=500):
        for i in range(0, len(lst), n):
            yield lst[i:i+n]

    bb = list(barcodes_to_load.values())
    for i, chunk in enumerate(chunks(bb)):
        sb.table('brand_barcodes').upsert(chunk, on_conflict='barcode').execute()
        print(f"  brand_barcodes: загружено {min((i+1)*500, len(bb))}/{len(bb)}")

    um = list(unmapped.values())
    for i, chunk in enumerate(chunks(um)):
        sb.table('unmapped_matrix_skus').upsert(chunk, on_conflict='barcode').execute()
        print(f"  unmapped_matrix_skus: загружено {min((i+1)*500, len(um))}/{len(um)}")

    print("\n=== Сверка ===")
    bb_cnt = sb.table('brand_barcodes').select('id', count='exact').execute().count
    um_cnt = sb.table('unmapped_matrix_skus').select('id', count='exact').execute().count
    print(f"  brand_barcodes: {bb_cnt}")
    print(f"  unmapped_matrix_skus: {um_cnt}")

if __name__ == '__main__':
    main()
