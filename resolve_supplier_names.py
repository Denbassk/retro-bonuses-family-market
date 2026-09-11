#!/usr/bin/env python3
"""
resolve_supplier_names.py - ЭТАП 2: сопоставление имён из Excel с suppliers.

Автоматически принимаются ТОЛЬКО:
  - записи, уже подтверждённые в retro_fact_name_map;
  - точные совпадения нормализованных имён.
Всё остальное -> CSV на ручное подтверждение. Скрипт ничего не пишет в БД.

Запуск:
  python resolve_supplier_names.py "Ретро Бонусы.xlsx" --sheet 2026
"""
import re, csv, argparse
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook

import sb
from import_retro_facts import norm_name, STOP_NAMES

NOISE = ("от оплат", "от оплаты", "от оплати", "кроме", "крім")


def split_name(nm):
    """Делит нормализованное имя на базу и содержимое скобок."""
    brands = " ".join(re.findall(r"\(([^)]*)\)", nm))
    base = re.sub(r"\([^)]*\)", " ", nm)
    for n in NOISE:
        base = base.replace(n, " ")
        brands = brands.replace(n, " ")
    tidy = lambda s: re.sub(r"\s+", " ", s).strip(" ,.-")
    return tidy(base), tidy(brands)


def score(ex, db):
    """0..1. База важнее, но совпадение бренда в скобках сильно поднимает."""
    eb, ebr = split_name(ex)
    db_, dbr = split_name(db)
    s = SequenceMatcher(None, eb, db_).ratio() * 0.55
    if ebr and dbr:
        r = SequenceMatcher(None, ebr, dbr).ratio()
        s += r * 0.45
        if ebr == dbr or ebr in dbr or dbr in ebr:
            s = min(1.0, s + 0.15)
    elif not ebr and not dbr:
        s += 0.45
    return round(s, 3)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    ws = load_workbook(Path(args.xlsx), data_only=True)[args.sheet]
    excel_names = []
    for r in range(args.header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or not str(v).strip():
            continue
        raw = str(v).strip()
        if norm_name(raw) in STOP_NAMES:
            continue
        excel_names.append((r, raw, norm_name(raw)))

    suppliers = sb.get("suppliers", "select=id,name,legal_name")
    sup_norm = {s["id"]: norm_name(s["name"]) for s in suppliers}
    by_norm = {}
    for s in suppliers:
        by_norm.setdefault(sup_norm[s["id"]], s)
    print(f"[i] Имён в Excel: {len(excel_names)} | поставщиков в Supabase: {len(suppliers)}")

    try:
        confirmed = {m["excel_name_normalized"]: m
                     for m in sb.get("retro_fact_name_map",
                                     "select=excel_name_normalized,supplier_id,is_ignored")}
    except RuntimeError as e:
        print(f"  [!] retro_fact_name_map недоступна ({e}); считаем пустой")
        confirmed = {}
    print(f"[i] Уже подтверждено: {len(confirmed)}")

    sup_name = {s["id"]: s["name"] for s in suppliers}
    rows, stat = [], {"CONFIRMED": 0, "IGNORED": 0, "EXACT": 0, "REVIEW": 0, "NO_CANDIDATE": 0}

    for r, raw, nm in excel_names:
        rec = {"row": r, "excel_name_raw": raw, "excel_name_norm": nm,
               "status": "", "supplier_id": "", "supplier_name": "", "score": "",
               "alt_1": "", "alt_2": "", "alt_3": "", "decision_supplier_id": ""}

        if nm in confirmed:
            c = confirmed[nm]
            if c.get("is_ignored"):
                rec["status"] = "IGNORED"
            else:
                rec["status"] = "CONFIRMED"
                rec["supplier_id"] = c["supplier_id"] or ""
                rec["supplier_name"] = sup_name.get(c["supplier_id"], "?")
        elif nm in by_norm:
            s = by_norm[nm]
            rec.update(status="EXACT", supplier_id=s["id"],
                       supplier_name=s["name"], score="1.0")
        else:
            ranked = sorted(((score(nm, sup_norm[s["id"]]), s) for s in suppliers),
                            key=lambda x: -x[0])[:4]
            if ranked and ranked[0][0] >= 0.45:
                rec["status"] = "REVIEW"
                rec["score"] = ranked[0][0]
                rec["supplier_name"] = ranked[0][1]["name"]
                for i, (sc, s) in enumerate(ranked[:3], 1):
                    rec[f"alt_{i}"] = f"{sc} | {s['name']} | {s['id']}"
            else:
                rec["status"] = "NO_CANDIDATE"

        stat[rec["status"]] += 1
        rows.append(rec)

    out = Path(args.out or f"name_review_{args.sheet}.csv")
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(rows)

    print("\n[i] Итог:")
    for k, v in stat.items():
        print(f"    {k:<13} {v}")
    print(f"\n[>] {out}")
    print("    Откройте в Excel. Для REVIEW/NO_CANDIDATE впишите нужный id")
    print("    в колонку decision_supplier_id (или IGNORE, если это не поставщик).")


if __name__ == "__main__":
    main()
