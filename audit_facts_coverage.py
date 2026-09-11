#!/usr/bin/env python3
"""
audit_facts_coverage.py - аудит покрытия фактов по всем месяцам.

Сводит три источника: ячейки Excel, retro_payments_fact (админка), retro_calculations (расчёт).
Ищет пропуски, конфликты и отличает перераспределение доплат от реальных расхождений
(по признаку: сумма дельт поставщика за год = 0).

Ничего не пишет в БД.

Запуск:
  python audit_facts_coverage.py "Ретро Бонусы.xlsx" --sheet 2026
"""
import re, csv, argparse
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import sb
from import_retro_facts import norm_name, parse_amount, classify_header, STOP_NAMES

CALC_COLS = ("total_retro_amount", "retro_amount_total", "retro_total",
             "total_amount", "amount_total", "retro_amount")
EPS = 0.5


def pick_calc_col(rows):
    if not rows:
        return None
    keys = rows[0].keys()
    for c in CALC_COLS:
        if c in keys:
            return c
    for k in keys:
        if "amount" in k and isinstance(rows[0][k], (int, float)):
            return k
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--upto", type=int, default=None, help="последний закрытый месяц")
    args = ap.parse_args()

    year = int(re.search(r"(20\d{2})", args.sheet).group(1))
    ws = load_workbook(Path(args.xlsx), data_only=True)[args.sheet]

    cols = {}
    for c in range(2, ws.max_column + 1):
        cl = classify_header(ws.cell(row=args.header_row, column=c).value, year)
        if cl and cl[0] == "month":
            cols[c] = cl[1]

    nmap = {m["excel_name_normalized"]: m for m in sb.get(
        "retro_fact_name_map",
        "select=excel_name_normalized,excel_name_raw,supplier_id,is_ignored,split_targets")}
    sup = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}

    excel = {}
    rowname = {}
    for r in range(args.header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or not str(v).strip():
            continue
        raw = str(v).strip()
        nm = norm_name(raw)
        if nm in STOP_NAMES:
            continue
        mp = nmap.get(nm)
        if not mp or mp.get("is_ignored") or mp.get("split_targets") or not mp.get("supplier_id"):
            continue
        sid = mp["supplier_id"]
        rowname[sid] = (r, raw)
        for c, per in cols.items():
            excel[(sid, per)] = (parse_amount(ws.cell(row=r, column=c).value),
                                 f"{get_column_letter(c)}{r}")

    facts = {}
    for x in sb.get("retro_payments_fact",
                    f"select=supplier_id,period_label,amount_paid&period_label=like.{year}-*"):
        facts[(x["supplier_id"], x["period_label"])] = float(x["amount_paid"])

    craw = sb.get("retro_calculations", f"select=*&period_label=like.{year}-*")
    ccol = pick_calc_col(craw)
    print(f"[i] Колонка суммы расчёта: {ccol or 'НЕ НАЙДЕНА'}")
    calc = {}
    if ccol:
        for x in craw:
            if x.get(ccol) is not None:
                calc[(x["supplier_id"], x["period_label"])] = float(x[ccol])

    months = sorted(cols.values())
    if args.upto:
        months = [m for m in months if int(m.split("-")[1]) <= args.upto]

    sids = {k[0] for k in excel} | {k[0] for k in facts} | {k[0] for k in calc}
    out, delta_by_sup, conflicts = [], defaultdict(float), defaultdict(list)
    stat = defaultdict(int)

    for sid in sids:
        in_sheet = sid in rowname
        for per in months:
            ex, cell = excel.get((sid, per), (None, ""))
            ft = facts.get((sid, per))
            cl = calc.get((sid, per))
            has_ex, has_ft, has_cl = ex is not None, ft is not None, cl is not None

            if not in_sheet:
                st = "NO_ROW_IN_EXCEL" if (has_ft or has_cl) else None
            elif has_ex and has_ft:
                d = ex - ft
                st = "MATCH" if abs(d) < EPS else "CONFLICT"
                if st == "CONFLICT":
                    delta_by_sup[sid] += d
                    conflicts[sid].append((per, cell, ft, ex, d))
            elif has_ex and not has_ft:
                st = "NOT_IMPORTED"
            elif has_ft and not has_ex:
                st = "FACT_ONLY"
            elif has_cl:
                st = "CALC_NO_FACT"
            else:
                st = None
            if not st:
                continue
            stat[st] += 1
            out.append({"supplier": sup.get(sid, sid), "period": per, "cell": cell,
                        "status": st, "excel": ex, "fact": ft, "calc": cl,
                        "delta_excel_fact": (ex - ft) if (has_ex and has_ft) else None})

    print(f"[i] Поставщиков: {len(sids)} | месяцев: {len(months)}\n")
    for k in ("MATCH", "CONFLICT", "NOT_IMPORTED", "FACT_ONLY", "CALC_NO_FACT", "NO_ROW_IN_EXCEL"):
        print(f"    {k:<17} {stat[k]}")

    redis = {s: d for s, d in delta_by_sup.items() if abs(d) < EPS}
    real = {s: d for s, d in delta_by_sup.items() if abs(d) >= EPS}

    print(f"\n[i] ПЕРЕРАСПРЕДЕЛЕНИЕ доплат (сумма дельт = 0) — {len(redis)} поставщиков.")
    print("    В админке доплата разнесена по месяцам, в Excel лежит в месяце платежа.")
    for s in redis:
        print(f"    {sup.get(s, s)[:40]}")
        for per, cell, ft, ex, d in sorted(conflicts[s]):
            print(f"        {per}  {cell:<6} админка {ft:>11,.0f} | excel {ex:>11,.0f} | {d:>+10,.0f}")

    print(f"\n[!] РЕАЛЬНЫЕ расхождения (сумма дельт != 0) — {len(real)} поставщиков:")
    for s, tot in sorted(real.items(), key=lambda x: -abs(x[1])):
        print(f"    {sup.get(s, s)[:40]}   итого {tot:>+12,.0f}")
        for per, cell, ft, ex, d in sorted(conflicts[s]):
            print(f"        {per}  {cell:<6} админка {ft:>11,.0f} | excel {ex:>11,.0f} | {d:>+10,.0f}")

    gaps = [r for r in out if r["status"] == "CALC_NO_FACT"]
    print(f"\n[!] Расчёт есть, оплаты нет ({len(gaps)}):")
    for r in sorted(gaps, key=lambda x: (-(x["calc"] or 0))):
        print(f"    {r['period']}  {r['supplier'][:38]:<38} начислено {r['calc']:>12,.2f}")

    fo = [r for r in out if r["status"] == "FACT_ONLY"]
    print(f"\n[~] В админке есть, в Excel пусто ({len(fo)}):")
    for r in sorted(fo, key=lambda x: x["period"]):
        print(f"    {r['period']}  {r['supplier'][:38]:<38} {r['fact']:>12,.0f}  (ячейка {r['cell']})")

    nr = sorted({r["supplier"] for r in out if r["status"] == "NO_ROW_IN_EXCEL"})
    print(f"\n[i] Нет строки в Excel — безнал/маркетинг ({len(nr)}):")
    for n in nr:
        print(f"    {n}")

    dst = Path(f"facts_coverage_{year}.csv")
    with dst.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(sorted(out, key=lambda x: (x["supplier"], x["period"])))
    print(f"\n[>] {dst}")


if __name__ == "__main__":
    main()
