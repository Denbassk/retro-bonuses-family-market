#!/usr/bin/env python3
"""
dump_cell_comments.py - вытаскивает примечания к ячейкам (красные треугольники).
Ничего не пишет в БД. Только CSV для изучения формата.

Запуск:
  python dump_cell_comments.py "Ретро Бонусы.xlsx" --sheet 2026
"""
import re, csv, zipfile, argparse
from pathlib import Path
from paths import OUT
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from import_retro_facts import norm_name, parse_amount, classify_header, STOP_NAMES

NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NS_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"

DATE_RE = re.compile(r"\b(\d{1,2})[.\-/](\d{1,2})(?:[.\-/](\d{2,4}))?\b")
NUM_RE = re.compile(r"\d[\d\s\u00a0]{2,}(?:[.,]\d{1,2})?")


def comments_from_zip(path, sheet_name):
    """Резервный путь: читаем примечания напрямую из xlsx-архива."""
    out = {}
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rid = None
        for sh in wb.iter(f"{NS_MAIN}sheet"):
            if sh.get("name") == sheet_name:
                rid = sh.get(f"{NS_REL}id")
        if not rid:
            return out
        target = None
        for rel in rels.iter(f"{NS_PKG}Relationship"):
            if rel.get("Id") == rid:
                target = rel.get("Target").lstrip("/")
        if not target:
            return out
        if not target.startswith("xl/"):
            target = "xl/" + target
        srel = target.rsplit("/", 1)
        srel = f"{srel[0]}/_rels/{srel[1]}.rels"
        if srel not in z.namelist():
            return out
        crels = ET.fromstring(z.read(srel))
        cpath = None
        for rel in crels.iter(f"{NS_PKG}Relationship"):
            t = rel.get("Target", "")
            if "comments" in t.lower():
                cpath = t.replace("../", "xl/").lstrip("/")
        if not cpath or cpath not in z.namelist():
            return out
        croot = ET.fromstring(z.read(cpath))
        for cm in croot.iter(f"{NS_MAIN}comment"):
            txt = "".join(t.text or "" for t in cm.iter(f"{NS_MAIN}t"))
            if txt.strip():
                out[cm.get("ref")] = txt
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    args = ap.parse_args()

    path = Path(args.xlsx)
    year = int(re.search(r"(20\d{2})", args.sheet).group(1))
    ws = load_workbook(path, data_only=True)[args.sheet]

    cols = {}
    for c in range(2, ws.max_column + 1):
        cl = classify_header(ws.cell(row=args.header_row, column=c).value, year)
        if cl and cl[0] != "skip":
            cols[c] = cl

    names = {}
    for r in range(args.header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and str(v).strip():
            names[r] = str(v).strip()

    rows = []
    for r in range(args.header_row + 1, ws.max_row + 1):
        for c in cols:
            cm = ws.cell(row=r, column=c).comment
            if cm and (cm.text or "").strip():
                rows.append((r, c, cm.text))

    if not rows:
        print("[i] openpyxl примечаний не отдал, читаю архив напрямую...")
        zc = comments_from_zip(path, args.sheet)
        for ref, txt in zc.items():
            m = re.match(r"([A-Z]+)(\d+)", ref)
            if not m:
                continue
            from openpyxl.utils import column_index_from_string
            c = column_index_from_string(m.group(1))
            r = int(m.group(2))
            if c in cols and r in names:
                rows.append((r, c, txt))

    out_rows = []
    for r, c, txt in sorted(rows):
        kind, label = cols[c]
        flat = re.sub(r"\s+", " ", txt).strip()
        dates = ["-".join(filter(None, (d[2] or str(year), d[1].zfill(2), d[0].zfill(2))))
                 for d in DATE_RE.findall(flat)]
        nums = [n.replace("\u00a0", "").replace(" ", "").replace(",", ".")
                for n in NUM_RE.findall(flat)]
        out_rows.append({
            "cell": f"{get_column_letter(c)}{r}",
            "supplier_raw": names.get(r, ""),
            "kind": kind,
            "period_label": label,
            "amount": parse_amount(ws.cell(row=r, column=c).value),
            "comment": flat,
            "dates_found": " | ".join(dates),
            "numbers_found": " | ".join(nums),
        })

    if not out_rows:
        print("[!] Примечаний не найдено вообще.")
        return

    out = OUT / f"cell_comments_{args.sheet}.csv"
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out_rows[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(out_rows)

    print(f"[i] Примечаний найдено: {len(out_rows)}")
    print("\n[i] Первые 12:")
    for x in out_rows[:12]:
        print(f"  {x['cell']:<6} {x['supplier_raw'][:28]:<28} {x['period_label']:<12} "
              f"сумма={x['amount']}")
        print(f"         {x['comment'][:110]}")
    print(f"\n[>] {out}")


if __name__ == "__main__":
    main()
