#!/usr/bin/env python3
"""
reconcile_facts.py - ЭТАП 4: автосверка факт оплат ретро vs расчёт, пара поставщик-месяц.

Источники: retro_payments_fact (админка), retro_calculations.total_retro (с доплатами и фикс-бонусом),
последний файл из Ретро_Excel\\ (для пометок «в Excel иначе» и составной строки БІР),
BigQuery loaded_at и supplier_payments.created_at (свежесть расчёта для незакрытых месяцев).

Правила (решения 11.09.2026, CLAUDE.md):
  - окно 2026-01 .. последний закрытый месяц; текущий месяц не сверяем;
  - янв-июн истина = сохранённый расчёт, свежесть не проверяем (TRUTH_UNTIL);
  - None (не заполнено) и 0 (платить нечего) - разные статусы;
  - оплата приходит в течение 2 мес. (95% фактов) -> пустой факт в этом окне = PENDING, позже = NOT_PAID;
  - безнал-поставщики в наличном Excel не ожидаются (NOT_EXPECTED);
  - составная строка (БІР Кег + Славутич) сверяется суммой по группе;
  - Оболонь: только цифры, без выводов (EXCLUDED_OWNER).
Пишет ТОЛЬКО в retro_reconciliation и только с --apply. Факты и расчёты не трогает.

Запуск (из корня):
  python core\\reconcile_facts.py                # показать + output\\reconciliation_2026.csv
  python core\\reconcile_facts.py --apply        # + записать в retro_reconciliation
  python core\\reconcile_facts.py --no-bq        # без проверки свежести по BigQuery
"""
import os, re, csv, uuid, argparse
from datetime import date, datetime, timezone, timedelta
from collections import defaultdict
from pathlib import Path

from paths import OUT, EXCEL_DIR
import sb
from openpyxl import load_workbook
from import_retro_facts import norm_name, parse_amount, classify_header, STOP_NAMES
from import_facts_to_db import spread_blocks

TRUTH_UNTIL = "2026-06"
LAG_MONTHS = 2
TOL_ABS = 1.0
TOL_PCT = 0.02            # ±2% = зелёная зона «Факт оплат ретро» в админке (getReconStatus)
REDIST_PCT = 0.005        # сумма расхождений за окно, при которой разнознаковые месяцы = перераспределение
NOT_EXPECTED_IN_EXCEL = {"Ново-Баварський", "Аванта-Трейд (Фереро, Киндер)", "Аванта-Трейд (АВК)",
                         "Еліт Фуд", "Сервіс Про"}
FIX_BONUS_NOTE = "фиксированный ежемесячный бонус"


def ts(v):
    """ISO-строка Supabase / datetime BigQuery -> aware datetime (UTC)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    m = re.match(r"(\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}:\d{2})(?:\.\d+)?(Z|[+-]\d{2}:?\d{2})?", str(v))
    t = datetime.fromisoformat(m.group(1).replace(" ", "T"))
    tz = (m.group(2) or "+00:00").replace("Z", "+00:00")
    if len(tz) == 5:
        tz = tz[:3] + ":" + tz[3:]
    return datetime.fromisoformat(t.isoformat() + tz).astimezone(timezone.utc)


def months(a, b):
    y, m = int(a[:4]), int(a[5:])
    out = []
    while f"{y}-{m:02d}" <= b:
        out.append(f"{y}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def month_add(per, k):
    y, m = int(per[:4]), int(per[5:]) + k
    y += (m - 1) // 12
    return f"{y}-{(m - 1) % 12 + 1:02d}"


def last_closed(today=None):
    t = today or date.today()
    return month_add(f"{t.year}-{t.month:02d}", -1)


def newest_excel():
    files = [p for p in EXCEL_DIR.glob("*.xlsx") if not p.name.startswith("~$")]
    return max(files, key=lambda p: p.stat().st_mtime) if files else None


def money(x):
    return "" if x is None else f"{x:,.0f}".replace(",", " ")


def load_excel(path, sheet, nmap, name2id):
    """-> {(key, per): (amount|None, cell)}, key = tuple(supplier_id...)"""
    from openpyxl.utils import get_column_letter
    year = int(re.search(r"(20\d{2})", sheet).group(1))
    ws = load_workbook(path, data_only=True)[sheet]
    cols = {}
    for c in range(2, ws.max_column + 1):
        cl = classify_header(ws.cell(row=1, column=c).value, year)
        if cl and cl[0] == "month":
            cols[c] = cl[1]
    out = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or not str(v).strip() or norm_name(str(v).strip()) in STOP_NAMES:
            continue
        mp = nmap.get(norm_name(str(v).strip()))
        if not mp or mp.get("is_ignored"):
            continue
        if mp.get("split_targets"):
            key = tuple(sorted(t.get("supplier_id") or name2id.get(t.get("supplier_name")) for t in mp["split_targets"]))
        elif mp.get("supplier_id"):
            key = (mp["supplier_id"],)
        else:
            continue
        for c, per in cols.items():
            a = parse_amount(ws.cell(row=r, column=c).value)
            if a is not None:
                out[(key, per)] = (a, f"{get_column_letter(c)}{r}")
    return out


def data_loaded_bq(per_from, alias2sids):
    """-> {(supplier_id, per): max loaded_at} по приходам и возвратам BQ начиная с per_from."""
    from google.cloud import bigquery
    proj = os.environ.get("BQ_PROJECT", "family-market-analytics")
    ds = os.environ.get("BQ_DATASET", "family_market")
    # loaded_at - DATETIME без пояса, записан по КИЕВСКОМУ времени (проверено 11.09: MAX = 14:16 при 13:56 UTC)
    sql = f"""SELECT supplier, per, MAX(ts) ts FROM (
        SELECT supplier, FORMAT_DATE('%Y-%m', doc_date) per, TIMESTAMP(loaded_at, 'Europe/Kiev') ts
        FROM `{proj}.{ds}.incoming_transactions` WHERE doc_date >= '{per_from}-01'
        UNION ALL
        SELECT supplier, FORMAT_DATE('%Y-%m', doc_date) per, TIMESTAMP(loaded_at, 'Europe/Kiev') ts
        FROM `{proj}.{ds}.outgoing_to_supplier_transactions` WHERE doc_date >= '{per_from}-01'
      ) GROUP BY 1, 2"""
    out = {}
    for r in bigquery.Client(project=proj).query(sql):
        for sid in alias2sids.get((r["supplier"] or "").strip(), ()):
            k = (sid, r["per"])
            out[k] = max(out.get(k, r["ts"]), r["ts"])
    return out


def data_loaded_payments(per_from):
    """-> {(supplier_id, per): max created_at} банковской выписки (для retro_base_type=payments)."""
    out = {}
    for p in sb.get("supplier_payments", f"period_label=gte.{per_from}&select=supplier_id,period_label,created_at"):
        if p.get("supplier_id") and p.get("created_at"):
            k = (p["supplier_id"], p["period_label"])
            t = ts(p["created_at"])
            out[k] = max(out.get(k, t), t)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=None, help="файл Excel (по умолчанию - самый новый в Ретро_Excel)")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--from", dest="per_from", default="2026-01")
    ap.add_argument("--to", dest="per_to", default=None, help="по умолчанию - последний закрытый месяц")
    ap.add_argument("--no-bq", action="store_true")
    ap.add_argument("--no-diagnose", action="store_true", help="не запускать диагностику после записи")
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    per_to = a.per_to or last_closed()
    pers = months(a.per_from, per_to)
    run_id = str(uuid.uuid4())

    sup = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}
    name2id = {v: k for k, v in sup.items()}
    nmap = {m["excel_name_normalized"]: m for m in sb.get(
        "retro_fact_name_map", "select=excel_name_normalized,supplier_id,is_ignored,split_targets,needs_manual")}
    manual_sids = {m["supplier_id"] for m in nmap.values() if m.get("needs_manual") and m.get("supplier_id")}
    groups = {}
    for m in nmap.values():
        if m.get("split_targets"):
            g = tuple(sorted(t.get("supplier_id") or name2id.get(t.get("supplier_name")) for t in m["split_targets"]))
            for sid in g:
                groups[sid] = g

    facts = defaultdict(list)
    for f in sb.get("retro_payments_fact", f"period_label=gte.{a.per_from}&period_label=lte.{per_to}"
                    "&select=id,supplier_id,period_label,amount_paid,payment_date,needs_manual,import_source"):
        facts[(f["supplier_id"], f["period_label"])].append(f)
    calcs = defaultdict(list)
    for c in sb.get("retro_calculations", f"period_label=gte.{a.per_from}&period_label=lte.{per_to}"
                    "&select=id,supplier_id,period_label,total_retro,status,created_at"):
        calcs[(c["supplier_id"], c["period_label"])].append(c)

    xlsx = Path(a.xlsx) if a.xlsx else newest_excel()
    excel = load_excel(xlsx, a.sheet, nmap, name2id) if xlsx else {}
    # доплаты «заплачено отдельно» (ДМП Славутич): в «Оплачено» и в расчёт не входят, но в Excel сумма общая
    separate = defaultdict(float)
    for adj in sb.get("retro_adjustments", f"period_label=gte.{a.per_from}&period_label=lte.{per_to}&select=*"):
        if (adj.get("payment_mode") or "in_payment") == "separate":
            separate[(adj["supplier_id"], adj["period_label"])] += float(adj["amount"] or 0)

    # ── свежесть расчёта (только незакрытые для «истины» месяцы) ──
    fresh_from = month_add(TRUTH_UNTIL, 1)
    alias2sids = defaultdict(set)
    for al in sb.get("supplier_aliases", "alias_type=neq.excluded&select=alias_name,supplier_id"):
        if al.get("alias_name") and al.get("supplier_id"):
            alias2sids[al["alias_name"].strip()].add(al["supplier_id"])
    loaded = {} if a.no_bq else data_loaded_bq(fresh_from, alias2sids)
    for k, t in data_loaded_payments(fresh_from).items():
        loaded[k] = max(loaded.get(k, t), t)

    this_month = month_add(per_to, 1)
    keys = {groups.get(s, (s,)) for s, _ in list(facts) + list(calcs)} | {k for k, _ in excel}
    rows, by_key = [], defaultdict(list)

    for key in keys:
        names = " + ".join(sup.get(s, "?") for s in key)
        # ручная разноска: Excel и админка по месяцам разные, нарастающим итогом равны (как в черновике загрузки)
        seq = []
        for per in pers:
            fr = [f for s in key for f in facts.get((s, per), [])]
            sp_ = sum(separate.get((s, per), 0.0) for s in key)
            adm = sum(float(f["amount_paid"]) for f in fr) + sp_ if fr or sp_ else None
            xa = excel.get((key, per), (None, ""))[0]
            seq.append((per, xa or 0.0, adm, (xa or 0.0) - (adm or 0.0)))
        spread = {}
        for i, j in spread_blocks(seq):
            for p_ in pers[i:j + 1]:
                spread[p_] = f"{pers[i]}..{pers[j]}"
        for per in pers:
            frows = [f for s in key for f in facts.get((s, per), [])]
            crows = [c for s in key for c in calcs.get((s, per), [])]
            fact = round(sum(float(f["amount_paid"]) for f in frows), 2) if frows else None
            calc = round(sum(float(c["total_retro"] or 0) for c in crows), 2) if crows else None
            ex_amt, ex_cell = excel.get((key, per), (None, ""))
            if fact is None and not calc and ex_amt is None:
                continue
            diag, dd = [], {}
            delta = round(fact - calc, 2) if fact is not None and calc is not None else None

            if any("оболон" in sup.get(s, "").lower() for s in key):
                st = "EXCLUDED_OWNER"
            elif fact is None:
                if ex_amt is not None:
                    st = "EXCEL_ONLY"
                    diag.append(f"в Excel {ex_cell} = {money(ex_amt)}, в админке пусто")
                    if calc:
                        diag.append(f"Excel - расчёт = {ex_amt - calc:+,.0f} ({(ex_amt - calc) / calc:+.1%})")
                elif all(sup.get(s) in NOT_EXPECTED_IN_EXCEL for s in key):
                    st = "NOT_EXPECTED"
                    diag.append("безнал: в наличном Excel не бывает, факт вносится только в админке")
                elif month_add(per, LAG_MONTHS) >= this_month:
                    st = "PENDING"
                else:
                    st = "NOT_PAID"
                    diag.append(f"начислено {money(calc)}, оплаты нет больше {LAG_MONTHS} мес.")
            elif abs(fact) < 0.005:
                st = "ZERO_OK" if not calc or abs(calc) < 0.005 else "ZERO_BUT_CALC"
                if st == "ZERO_BUT_CALC":
                    diag.append(f"внесён 0, а начислено {money(calc)}")
            elif not calc:
                st = "FACT_NO_CALC"
                diag.append(f"оплачено {money(fact)}, расчёта нет" if calc is None else f"оплачено {money(fact)}, расчёт = 0")
            elif abs(delta) < TOL_ABS:
                st = "MATCH"
            elif abs(delta) <= TOL_PCT * abs(calc):
                st = "MINOR"
            else:
                st = "MISMATCH"

            # пометки, не меняющие статус
            sep = round(sum(separate.get((s, per), 0.0) for s in key), 2)
            if sep:
                dd["separate_extras"] = sep
            if ex_amt is not None and fact is not None and abs(ex_amt - fact - sep) >= TOL_ABS:
                if per in spread:
                    diag.append(f"Excel и админка разнесены по месяцам по-разному, итог за {spread[per]} совпадает")
                    dd["manual_spread"] = spread[per]
                else:
                    diag.append(f"в Excel {ex_cell} = {money(ex_amt)}, в админке {money(fact)}"
                                + (f" + отдельно {money(sep)}" if sep else ""))
                    dd["excel_vs_admin"] = round(ex_amt - fact - sep, 2)
            y, m = int(per[:4]), int(per[5:])
            for f in frows:
                pdt = f.get("payment_date")
                try:
                    ok = pdt is None or -31 <= (date.fromisoformat(pdt) - date(y, m, 1)).days <= 240
                except ValueError:
                    ok = False
                if not ok:
                    diag.append(f"дата оплаты «{pdt}» подозрительна")
                    dd["bad_payment_date"] = pdt
            stale = []
            for c in crows:
                lt = loaded.get((c["supplier_id"], per))
                if per > TRUTH_UNTIL and c["status"] not in ("approved", "paid") and lt and lt > ts(c["created_at"]):
                    stale.append((sup.get(c["supplier_id"]), ts(c["created_at"]), lt))
            if stale:
                dd["stale"] = [{"supplier": n, "calc": str(ct), "data": str(lt)} for n, ct, lt in stale]

            r = {"key": key, "names": names, "per": per, "fact": fact, "calc": calc, "delta": delta,
                 "ex_amt": ex_amt, "ex_cell": ex_cell, "status": st, "diag": diag, "dd": dd, "stale": stale,
                 "frows": frows, "crows": crows}
            rows.append(r)
            by_key[key].append(r)

    # ── уточнение расхождений: устаревший расчёт -> перераспределение -> ручной поставщик ──
    for key, rs in by_key.items():
        both = [r for r in rs if r["delta"] is not None]
        s_delta = round(sum(r["delta"] for r in both), 2)
        s_calc = sum(abs(r["calc"]) for r in both)
        # перераспределение = разнознаковые расхождения, гасящиеся в сумме до 0,5% (не 2%: иначе
        # систематический недосчёт +2..3% каждый месяц, как у Бісквіт і Ко, выглядел бы «разноской»)
        sig = [r["delta"] for r in both if abs(r["delta"]) >= TOL_ABS]
        redistributed = (len(both) >= 2 and any(d > 0 for d in sig) and any(d < 0 for d in sig)
                         and abs(s_delta) < max(TOL_ABS * len(both), REDIST_PCT * s_calc))
        for r in rs:
            if r["status"] != "MISMATCH":
                continue
            if r["stale"]:
                n, ct, lt = r["stale"][0]
                kyiv = timezone(timedelta(hours=3))   # EEST; для зимы - UTC+2, на вывод не влияет по сути
                r["status"] = "STALE_CALC"
                r["diag"].insert(0, f"расчёт от {ct.astimezone(kyiv):%d.%m %H:%M}, данные перезалиты "
                                    f"{lt.astimezone(kyiv):%d.%m %H:%M} (Киев) - "
                                    f"сначала пересчитать ({n})")
            elif redistributed:
                r["status"] = "REDISTRIBUTION"
                r["diag"].insert(0, f"за {both[0]['per']}..{both[-1]['per']} факт-расчёт в сумме {s_delta:+,.0f}: "
                                    f"оплата разнесена по другим месяцам")
                r["dd"]["window_delta"] = s_delta
            elif any(s in manual_sids for s in key):
                r["status"] = "MANUAL"
                r["diag"].insert(0, "поставщик помечен как ручной (компенсации акций в сумме)")
            if r["status"] in ("MISMATCH", "MANUAL"):
                r["diag"].insert(0, f"расхождение {r['delta']:+,.2f} ({r['delta'] / r['calc']:+.1%})")

    # ── вывод ──
    ORDER = ["MATCH", "MINOR", "ZERO_OK", "REDISTRIBUTION", "MISMATCH", "MANUAL", "STALE_CALC", "ZERO_BUT_CALC",
             "FACT_NO_CALC", "EXCEL_ONLY", "NOT_PAID", "PENDING", "NOT_EXPECTED", "EXCLUDED_OWNER"]
    cnt = defaultdict(lambda: defaultdict(int))
    for r in rows:
        cnt[r["per"]][r["status"]] += 1
    used = [s for s in ORDER if any(cnt[p][s] for p in pers)]
    print(f"[i] Сверка {a.per_from}..{per_to} | Excel: {xlsx.name if xlsx else '-'} | "
          f"BQ: {'выкл' if a.no_bq else 'да'} | истина до {TRUTH_UNTIL} = сохранённый расчёт\n")
    print(f"    {'период':<9}" + "".join(f"{s[:10]:>11}" for s in used))
    for p in pers:
        print(f"    {p:<9}" + "".join(f"{cnt[p][s] or '':>11}" for s in used))

    attention = ["MISMATCH", "MANUAL", "STALE_CALC", "ZERO_BUT_CALC", "FACT_NO_CALC", "EXCEL_ONLY", "NOT_PAID", "REDISTRIBUTION"]
    for s in attention:
        rs = sorted([r for r in rows if r["status"] == s], key=lambda r: (r["per"], r["names"]))
        if not rs:
            continue
        print(f"\n[{s}] {len(rs)}")
        for r in rs:
            print(f"    {r['per']}  {r['names'][:44]:<45} факт {money(r['fact']):>9}  расчёт {money(r['calc']):>9}"
                  f"  | {'; '.join(r['diag'])}")
    other = [r for r in rows if r["diag"] and r["status"] not in attention]
    if other:
        quiet = [r for r in other if r["status"] in ("MATCH", "MINOR", "NOT_EXPECTED")]
        print(f"\n[пометки] {len(other)}: из них {len(quiet)} при совпавшем расчёте или безнал "
              f"(Excel держит оплату в месяце платежа, админка - в месяце начисления) - в CSV")
        for r in sorted([r for r in other if r not in quiet], key=lambda r: (r["per"], r["names"])):
            print(f"    {r['per']}  {r['names'][:44]:<45} {r['status']:<10} | {'; '.join(r['diag'])}")
    stale_quiet = [r for r in rows if r["stale"] and r["status"] != "STALE_CALC"]
    if stale_quiet:
        print(f"[i] расчёт старше данных, но сходится или ждёт оплаты: {len(stale_quiet)} пар (в diagnosis_data)")

    dst = OUT / f"reconciliation_{a.sheet}.csv"
    with dst.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["period", "supplier", "status", "fact", "calc", "delta", "excel", "excel_cell", "diagnosis"])
        for r in sorted(rows, key=lambda r: (r["per"], r["names"])):
            w.writerow([r["per"], r["names"], r["status"], r["fact"], r["calc"], r["delta"], r["ex_amt"],
                        r["ex_cell"], "; ".join(r["diag"])])
    print(f"\n[>] {dst.relative_to(OUT.parent)}")

    if not a.apply:
        print("[i] Пробный прогон. Для записи в retro_reconciliation добавьте --apply")
        return

    now = datetime.now(timezone.utc).isoformat()
    batch = []
    for r in rows:
        group = r["names"] if len(r["key"]) > 1 else None
        for sid in r["key"]:
            fr = [f for f in r["frows"] if f["supplier_id"] == sid]
            cr = [c for c in r["crows"] if c["supplier_id"] == sid]
            lt = loaded.get((sid, r["per"]))
            dd = dict(r["dd"])
            if group:
                dd.update(group_fact=r["fact"], group_calc=r["calc"])
            batch.append({
                "supplier_id": sid, "period_label": r["per"], "group_key": group,
                "fact_amount": round(sum(float(f["amount_paid"]) for f in fr), 2) if fr else None,
                "excel_amount": r["ex_amt"], "excel_cell": r["ex_cell"] or None,
                "calc_amount": round(sum(float(c["total_retro"] or 0) for c in cr), 2) if cr else None,
                "calc_status": cr[0]["status"] if cr else None,
                "calc_created_at": cr[0]["created_at"] if cr else None,
                "data_loaded_at": lt.isoformat() if lt else None,
                "delta": r["delta"],
                "delta_pct": round(r["delta"] / r["calc"] * 100, 4) if r["delta"] is not None and r["calc"] else None,
                "status": r["status"], "diagnosis": "; ".join(r["diag"]) or None, "diagnosis_data": dd or None,
                "run_id": run_id, "checked_at": now})
    sb.upsert("retro_reconciliation", batch, on_conflict="supplier_id,period_label")
    # пары, которые в этом прогоне не появились (факт/расчёт удалены), убираем - только внутри окна
    sb._req(f"{sb.URL}/rest/v1/retro_reconciliation?period_label=gte.{a.per_from}&period_label=lte.{per_to}"
            f"&run_id=neq.{run_id}", None, "DELETE", {"Prefer": "return=minimal"})
    print(f"[>] retro_reconciliation: записано {len(batch)} (run {run_id[:8]})")

    if not a.no_diagnose:
        import diagnose_retro   # импорт здесь: diagnose_retro сам импортирует этот модуль
        print("\n[i] Этап 5: диагностика проблемных пар (BigQuery, эталоны, правила, примечания)...")
        res = diagnose_retro.run(apply=True, no_bq=a.no_bq, verbose=False)
        closed = sum(1 for _, r in res if r["closed"])
        print(f"[i] объяснено полностью {closed} из {sum(1 for _, r in res if r['closed'] is not None)} расхождений")
        if not a.no_bq:
            import bq_docs
            print("\n[i] Полнота базы: документы BigQuery против эталона Торгсофт (все поставщики)...")
            bq_docs.health(a.per_from, per_to)


if __name__ == "__main__":
    main()
