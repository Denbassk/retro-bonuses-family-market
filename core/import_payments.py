#!/usr/bin/env python3
"""
import_payments.py — импорт банковских выписок (Excel) в supplier_payments.

Структура Excel:
  Строки 1-5: служебные (название организации, заголовок, фильтры, пусто, пусто)
  Строка 6: заголовки столбцов
  Строки 7..N: данные
  Последние строки: "Разом", "Відповідальний:" и подписи — пропускаем.

Маппинг поставщика — по подстроке legal_name в "Інформація".
"""

import os
import sys
import json
import urllib.request
import urllib.error
from pathlib import Path
from datetime import datetime
from decimal import Decimal

try:
    from openpyxl import load_workbook
except ImportError:
    print("ОШИБКА: установите openpyxl: pip install openpyxl")
    sys.exit(1)

# ─── Загрузка .env ─────────────────────────────────────────────────────────
def load_env(env_path=None):
    if env_path is None:
        # Ищем .env в текущей папке и в родительской
        here = Path(__file__).parent
        for candidate in [here / ".env", here.parent / ".env"]:
            if candidate.exists():
                env_path = candidate
                break
    if not env_path or not Path(env_path).exists():
        return
    with open(env_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            key, _, val = line.partition("=")
            os.environ.setdefault(key.strip(), val.strip())
load_env()
print("DEBUG ENV:", "SUPABASE_URL" in os.environ, list(os.environ.keys())[-5:])

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_KEY = os.environ["SUPABASE_SERVICE_KEY"]

# ─── Supabase helpers ──────────────────────────────────────────────────────
def sb_get(table, params=""):
    url = f"{SUPABASE_URL}/rest/v1/{table}?{params}"
    req = urllib.request.Request(url, headers={
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
    })
    with urllib.request.urlopen(req) as resp:
        return json.loads(resp.read())

def sb_post(table, data, on_conflict=None):
    url = f"{SUPABASE_URL}/rest/v1/{table}"
    if on_conflict:
        url += f"?on_conflict={on_conflict}"
    body = json.dumps(data, default=str).encode()
    headers = {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "return=representation,resolution=ignore-duplicates",
    }
    req = urllib.request.Request(url, data=body, method="POST", headers=headers)
    try:
        with urllib.request.urlopen(req) as resp:
            txt = resp.read()
            return json.loads(txt) if txt else []
    except urllib.error.HTTPError as e:
        err = e.read().decode("utf-8", errors="replace")
        print(f"[!] Supabase POST {table} error {e.code}: {err}")
        return None

# ─── Парсинг даты ──────────────────────────────────────────────────────────
def parse_date(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if hasattr(val, "year"):  # datetime.date
        return val
    s = str(val).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None

def parse_amount(val):
    if val is None or val == "":
        return None
    try:
        return float(str(val).replace(",", ".").replace(" ", ""))
    except ValueError:
        return None

# ─── Парсинг "Інформація" ──────────────────────────────────────────────────
import re

INVOICE_RE = re.compile(r"НАКЛАДНА\s*№\s*(\S+?)\s+ВIД\s+(\d{2}\.\d{2}\.\d{4})", re.IGNORECASE)
COUNTERPARTY_RE = re.compile(r'^([^/]+?)\s*/')

def parse_info(info_text):
    """Возвращает (counterparty, invoice_number, invoice_date)."""
    if not info_text:
        return None, None, None
    counterparty = None
    cp_match = COUNTERPARTY_RE.match(info_text)
    if cp_match:
        counterparty = cp_match.group(1).strip()
    invoice_number = None
    invoice_date = None
    inv_match = INVOICE_RE.search(info_text)
    if inv_match:
        invoice_number = inv_match.group(1).strip()
        invoice_date = parse_date(inv_match.group(2))
    return counterparty, invoice_number, invoice_date

# ─── Поиск поставщика по counterparty ──────────────────────────────────────
def normalize(s):
    if not s:
        return ""
    s = re.sub(r'["«»\']', '', s)
    s = re.sub(r'\s+', ' ', s)
    return s.strip().upper()


# Кеш маппинга — заполняется один раз при старте
_counterparty_map = None

def load_counterparty_map(suppliers_list):
    """Загружает таблицу bank_counterparty_map в кеш."""
    global _counterparty_map
    if _counterparty_map is not None:
        return _counterparty_map
    try:
        rows = sb_get("bank_counterparty_map", "select=counterparty_normalized,supplier_id")
    except Exception as e:
        print(f"  [!] Не удалось загрузить bank_counterparty_map: {e}")
        rows = []
    sup_by_id = {s["id"]: s["name"] for s in suppliers_list}
    _counterparty_map = {
        r["counterparty_normalized"]: (r["supplier_id"], sup_by_id.get(r["supplier_id"], "???"))
        for r in rows
    }
    print(f"[i] Загружено маппингов контрагентов: {len(_counterparty_map)}")
    return _counterparty_map


def find_supplier(counterparty, suppliers_list):
    """
    Ищет supplier_id по counterparty.
    Приоритет:
      1. Точное совпадение в bank_counterparty_map.
      2. Точное совпадение по legal_name.
      3. Частичное совпадение по legal_name.
    """
    if not counterparty:
        return None, None

    norm_cp = normalize(counterparty)

    # 1. Маппинг из bank_counterparty_map
    mapping = load_counterparty_map(suppliers_list)
    if norm_cp in mapping:
        return mapping[norm_cp]

    # 2. Точное по legal_name
    for s in suppliers_list:
        legal = s.get("legal_name") or ""
        if legal and normalize(legal) == norm_cp:
            return s["id"], s["name"]

    # 3. Частичное по legal_name
    for s in suppliers_list:
        legal = s.get("legal_name") or ""
        if legal and normalize(legal) in norm_cp:
            return s["id"], s["name"]

    return None, None

# ─── Импорт одного файла ───────────────────────────────────────────────────
def import_file(file_path, suppliers_list):
    print(f"\n[i] Файл: {file_path.name}")
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    # Найдём строку заголовков (ищем "№ з/п")
    header_row = None
    for r in range(1, 15):
        v = ws.cell(row=r, column=1).value
        if v and "з/п" in str(v).lower():
            header_row = r
            break
    if header_row is None:
        print(f"  [!] Не найдена строка заголовков")
        return 0, 0

    # Определим столбцы по заголовкам
    col_map = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=c).value
        if not h:
            continue
        h_norm = str(h).strip().lower()
        if "з/п" in h_norm: col_map["num"] = c
        elif h_norm == "дата": col_map["date"] = c
        elif "сума" in h_norm: col_map["amount"] = c
        elif "номер вх" in h_norm: col_map["invoice_num"] = c
        elif h_norm == "номер": col_map["doc_num"] = c
        elif "інформація" in h_norm or "информация" in h_norm: col_map["info"] = c

    required = ["date", "amount", "info"]
    if not all(k in col_map for k in required):
        print(f"  [!] Не хватает столбцов: {required} / нашли: {list(col_map.keys())}")
        return 0, 0

    saved = 0
    skipped = 0
    seen_supplier = None

    for r in range(header_row + 1, ws.max_row + 1):
        num_val = ws.cell(row=r, column=col_map["num"]).value
        # Стоп на "Разом" / "Відповідальний:"
        if num_val and isinstance(num_val, str):
            low = num_val.strip().lower()
            if low.startswith("разом") or low.startswith("відповідальний") or low.startswith("ответственный"):
                break

        date_val = parse_date(ws.cell(row=r, column=col_map["date"]).value)
        amount = parse_amount(ws.cell(row=r, column=col_map["amount"]).value)
        info = ws.cell(row=r, column=col_map["info"]).value
        if not date_val or amount is None or not info:
            continue

        doc_num = ws.cell(row=r, column=col_map["doc_num"]).value if "doc_num" in col_map else None
        invoice_num = ws.cell(row=r, column=col_map["invoice_num"]).value if "invoice_num" in col_map else None

        counterparty, inv_n, inv_d = parse_info(str(info))
        sup_id, sup_name = find_supplier(counterparty, suppliers_list)
        if not sup_id:
            if seen_supplier != counterparty:
                print(f"  [!] Не найден поставщик для: {counterparty}")
                seen_supplier = counterparty
            skipped += 1
            continue

        if seen_supplier != sup_name:
            print(f"  → {sup_name}")
            seen_supplier = sup_name

        period_label = f"{date_val.year}-{date_val.month:02d}"

        payload = {
            "supplier_id": sup_id,
            "payment_date": str(date_val),
            "amount": amount,
            "doc_number": str(int(doc_num)) if isinstance(doc_num, (int, float)) else (str(doc_num) if doc_num else None),
            "invoice_number": str(int(invoice_num)) if isinstance(invoice_num, (int, float)) else (str(invoice_num) if invoice_num else (inv_n or None)),
            "invoice_date": str(inv_d) if inv_d else None,
            "counterparty": counterparty,
            "raw_info": str(info)[:500],
            "source_file": file_path.name,
            "period_label": period_label,
        }
        res = sb_post("supplier_payments", payload,
                      on_conflict="supplier_id,payment_date,doc_number,amount")
        if res is None:
            skipped += 1
        else:
            saved += 1

    print(f"  Импортировано: {saved}, пропущено: {skipped}")
    return saved, skipped

# ─── main ──────────────────────────────────────────────────────────────────
def main():
    import argparse
    parser = argparse.ArgumentParser(description="Импорт банковских выписок в supplier_payments")
    parser.add_argument("--all", action="store_true",
                        help="Загрузить все файлы из папки payments/ (без диалога)")
    parser.add_argument("--files", nargs="+",
                        help="Конкретные файлы через пробел")
    args = parser.parse_args()

    payments_dir = Path(__file__).resolve().parent.parent / "payments"
    if not payments_dir.exists():
        payments_dir.mkdir()
        print(f"Создана папка: {payments_dir}")

    # Определяем список файлов
    files = []

    if args.files:
        # Режим: явно переданные файлы
        for f in args.files:
            p = Path(f)
            if not p.is_absolute():
                p = payments_dir / p
            if p.exists() and p.suffix.lower() == ".xlsx":
                files.append(p)
            else:
                print(f"[!] Пропущен (не найден): {f}")

    elif args.all:
        # Режим: все файлы из папки
        files = [f for f in payments_dir.glob("*.xlsx") if not f.name.startswith("~$")]

    else:
        # Режим по умолчанию: диалог выбора файлов
        try:
            import tkinter as tk
            from tkinter import filedialog
        except ImportError:
            print("[!] tkinter недоступен. Используйте --all или --files")
            sys.exit(1)

        root = tk.Tk()
        root.withdraw()  # скрываем основное окно
        root.attributes("-topmost", True)  # диалог поверх остальных окон

        selected = filedialog.askopenfilenames(
            title="Выберите Excel-файлы выписок для импорта",
            initialdir=str(payments_dir),
            filetypes=[("Excel файлы", "*.xlsx"), ("Все файлы", "*.*")],
        )
        root.destroy()

        if not selected:
            print("[i] Файлы не выбраны, выход.")
            sys.exit(0)

        files = [Path(f) for f in selected if Path(f).suffix.lower() == ".xlsx"]

    if not files:
        print("[!] Нет файлов для импорта.")
        sys.exit(1)

    print(f"\n[i] Выбрано файлов: {len(files)}")
    for f in files:
        print(f"    - {f.name}")
    print()

    suppliers_list = sb_get("suppliers", "select=id,name,legal_name")
    print(f"[i] Поставщиков в БД: {len(suppliers_list)}")

    total_saved = 0
    total_skipped = 0
    for f in files:
        s, sk = import_file(f, suppliers_list)
        total_saved += s
        total_skipped += sk

    print(f"\n{'='*60}")
    print(f"  ИТОГО: импортировано {total_saved}, пропущено {total_skipped}")
    print(f"{'='*60}")

    # Пауза, чтобы окно не закрылось сразу при двойном клике
    if sys.platform == "win32" and not (args.all or args.files):
        input("\nНажмите Enter для выхода...")


if __name__ == "__main__":
    main()
