#!/usr/bin/env python3
"""
diagnose_retro.py - ЭТАП 5: сам ищет причины расхождений факт vs расчёт (и недоплат, и переплат)
и причины отсутствия данных. Берёт проблемные пары из retro_reconciliation и проверяет лестницу гипотез.
Каждая гипотеза возвращает ЧИСЛО (сколько ₴ ретро она объясняет). Сильные принимаются, если не перебивают
остаток; слабые - только если закрывают разрыв целиком. Стоп, когда остаток в допуске.

Лестница (от дешёвых к дорогим):
  NOTE       примечание к ячейке Excel: нерет­ро компоненты (ДМП, кубы, стойки...), не внесённые доплатой
  EXCEL      в админке внесено иначе, чем в Excel, а Excel совпадает с расчётом
  SHIFT      разница гасится соседним месяцем (оплата перенесена)
  THRESHOLD  порог закупки (Монжар 80 тыс.) пройден/не пройден
  VAT        форма НДС: не вычтен (×1/0.8) или ÷1.2 вместо ×0.8
  RATE_XLS   ставка из колонки «Условия» Excel ≠ ставка правила
  RETURNS    возвраты: поставщик не вычел (subtract) / вычел при политике none   [BigQuery]
  SKU        приход по SKU вне правил, двойной счёт SKU, дрейф данных после расчёта   [BigQuery]
  ETALON     эталон Торгсофт: приход/возвраты больше или меньше наших   [BigQuery]
  PAYMENTS   база от оплат: выписка по дате платежа vs период   [Supabase]
  RATE_NICE  подразумеваемая «круглая» ставка (слабая)
  BOUNDARY   документы последних 3 дней месяца (слабая)   [BigQuery]
Для NOT_PAID / FACT_NO_CALC / ZERO_BUT_CALC / EXCEL_ONLY - поиск отсутствующих данных:
  последняя оплата, пустая ячейка, приход по алиасам в BQ, приход в эталоне, выписка, активные правила.

Запуск (из корня):
  python core\\diagnose_retro.py                          # все проблемные пары, отчёт + output\\diagnosis_2026.csv
  python core\\diagnose_retro.py --supplier "Злагода" --period 2026-07
  python core\\diagnose_retro.py --apply                  # записать диагноз в retro_reconciliation
Обычно запускается сам из reconcile_facts.py --apply.
"""
import os, re, argparse
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from paths import OUT
import sb
from openpyxl import load_workbook
from import_facts_to_db import note_body, NON_RETRO
from parse_note_components import components
from dump_cell_comments import comments_from_zip
from reconcile_facts import newest_excel, month_add, TOL_ABS, TOL_PCT

TARGET = ("MISMATCH", "MANUAL", "STALE_CALC", "NOT_PAID", "EXCEL_ONLY", "ZERO_BUT_CALC", "FACT_NO_CALC")
CLOSE_PCT = 0.005            # остаток <= 0,5% расчёта (не меньше 1 ₴) = разрыв закрыт
PROJ = os.environ.get("BQ_PROJECT", "family-market-analytics")
DS = os.environ.get("BQ_DATASET", "family_market")
FIX_BONUS = "фиксированный ежемесячный бонус"
MARK = "ДИАГНОЗ:"


def f2(x):
    return float(x or 0)


def money(x):
    return f"{x:+,.0f}"


def sp(text):
    """1,234,567 -> 1 234 567 только внутри чисел (запятые в названиях не трогаем)."""
    return re.sub(r"(?<=\d),(?=\d{3}(?!\d))", " ", text)


def bounds(per):
    y, m = int(per[:4]), int(per[5:])
    a = date(y, m, 1)
    return a, (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1))


def parse_rates(v):
    """«Условия» Excel: 0.15 | 15 | 'от 9,5-15%' -> [15.0] / [9.5, 15.0]"""
    if v is None:
        return []
    if isinstance(v, (int, float)):
        x = float(v)
        return [x * 100 if x < 1 else x] if 0 < x <= 50 else []
    out = []
    for s in re.findall(r"\d+(?:[.,]\d+)?", str(v)):
        x = float(s.replace(",", "."))
        x = x * 100 if x < 1 else x
        if 0 < x <= 50:
            out.append(x)
    return out


def in_chunks(table, col, ids, select, extra="", n=40):
    ids = list(ids)
    out = []
    for i in range(0, len(ids), n):
        out += sb.get(table, f"{col}=in.({','.join(ids[i:i + n])}){extra}&select={select}")
    return out


class Data:
    """Всё, что нужно гипотезам: справочники, расчёты, SKU-разбивка, факты, Excel, BigQuery."""

    def __init__(self, recon_all, targets, no_bq=False):
        self.sup = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name,returns_cutoff_day")}
        bsup = {b["id"]: b["supplier_id"] for b in sb.get("supplier_brands", "select=id,supplier_id")}
        self.rules = {}
        self.rules_by_sup = defaultdict(list)
        for r in sb.get("retro_rules", "status=eq.active&select=*"):
            r["supplier_id"] = bsup.get(r["supplier_brand_id"])
            self.rules[r["id"]] = r
            self.rules_by_sup[r["supplier_id"]].append(r)
        self.aliases = defaultdict(lambda: {"in": set(), "ret": set()})
        self.alias_owners = defaultdict(set)
        for a in sb.get("supplier_aliases", "alias_type=neq.excluded&select=alias_name,supplier_id,alias_type"):
            nm = (a.get("alias_name") or "").strip()
            if not nm or not a.get("supplier_id"):
                continue
            self.aliases[a["supplier_id"]]["ret"].add(nm)
            if a.get("alias_type") != "returns_only":
                self.aliases[a["supplier_id"]]["in"].add(nm)
                self.alias_owners[nm].add(a["supplier_id"])

        self.recon = {(r["supplier_id"], r["period_label"]): r for r in recon_all}
        pers = sorted({t["per"] for t in targets})
        self.pers = pers
        wide = sorted(set(pers) | {month_add(p, -1) for p in pers} | {month_add(p, 1) for p in pers})

        calcs = sb.get("retro_calculations", f"period_label=gte.{wide[0]}&period_label=lte.{wide[-1]}"
                       "&select=id,supplier_id,period_label,total_retro,total_purchase,total_returns,status,created_at")
        self.calc = {(c["supplier_id"], c["period_label"]): c for c in calcs}
        tcalc = [c["id"] for c in calcs if c["period_label"] in pers]
        self.details = defaultdict(list)
        for d in in_chunks("retro_calculation_details", "calculation_id", tcalc,
                           "id,calculation_id,retro_rule_id,amount_purchased,amount_returned,amount_net,"
                           "applied_percent,retro_amount,notes,adjustment_id"):
            self.details[d["calculation_id"]].append(d)
        # SKU, учтённые расчётом, по ВСЕМ поставщикам месяца (общие алиасы: чужой SKU не «вне правил»)
        self.sku_by = defaultdict(lambda: defaultdict(list))      # per -> barcode -> [(sid, detail_id, amount)]
        cid2 = {c["id"]: c for c in calcs}
        for s in in_chunks("retro_calculation_sku_details", "calculation_id", tcalc,
                           "calculation_id,detail_id,barcode,amount_purchased"):
            c = cid2[s["calculation_id"]]
            self.sku_by[c["period_label"]][str(s["barcode"])].append(
                (c["supplier_id"], s["detail_id"], f2(s["amount_purchased"])))

        sids = {s for t in targets for s in t["members"]}
        self.facts = {}
        for f in sb.get("retro_payments_fact", f"period_label=gte.{wide[0]}&period_label=lte.{wide[-1]}&select=*"):
            self.facts[(f["supplier_id"], f["period_label"])] = f
        self.facts_all = defaultdict(list)
        for f in in_chunks("retro_payments_fact", "supplier_id", sids, "supplier_id,period_label,amount_paid,payment_date"):
            self.facts_all[f["supplier_id"]].append(f)
        self.adjs = defaultdict(list)
        for a in in_chunks("retro_adjustments", "supplier_id", sids, "supplier_id,period_label,amount,notes"):
            self.adjs[(a["supplier_id"], a["period_label"])].append(a)
        self.payments = defaultdict(list)
        pay_sids = [s for s in sids if any(r.get("retro_base_type") == "payments" for r in self.rules_by_sup[s])]
        if pay_sids:
            for p in in_chunks("supplier_payments", "supplier_id", pay_sids, "supplier_id,payment_date,period_label,amount"):
                self.payments[p["supplier_id"]].append(p)

        self.load_excel()
        self.bq = {} if no_bq else self.load_bq(targets, wide)

    def load_excel(self):
        self.xlsx = newest_excel()
        self.comments, self.conditions = {}, {}
        if not self.xlsx:
            return
        ws = load_workbook(self.xlsx, data_only=True)["2026"]
        for r in range(2, ws.max_row + 1):
            self.conditions[r] = ws.cell(row=r, column=2).value
            for c in range(3, ws.max_column + 1):
                cm = ws.cell(row=r, column=c).comment
                if cm and (cm.text or "").strip():
                    self.comments[f"{ws.cell(row=r, column=c).column_letter}{r}"] = cm.text
        if not self.comments:
            self.comments = comments_from_zip(self.xlsx, "2026")

    def load_bq(self, targets, wide):
        from google.cloud import bigquery
        names_in = sorted({n for t in targets for s in t["members"] for n in self.aliases[s]["in"]})
        names_ret = sorted({n for t in targets for s in t["members"] for n in self.aliases[s]["ret"]})
        if not names_in and not names_ret:
            return {}
        d1, d2 = bounds(wide[0])[0], bounds(wide[-1])[1]
        cli = bigquery.Client(project=PROJ)

        def run(sql, names):
            cfg = bigquery.QueryJobConfig(query_parameters=[
                bigquery.ArrayQueryParameter("n", "STRING", names or ["-"]),
                bigquery.ScalarQueryParameter("d1", "DATE", d1), bigquery.ScalarQueryParameter("d2", "DATE", d2)])
            return list(cli.query(sql, job_config=cfg))

        T = lambda t: f"`{PROJ}.{DS}.{t}`"
        out = {"inc": defaultdict(float), "inc_last3": defaultdict(float), "raw": defaultdict(float),
               "ret": defaultdict(float), "et_inc": defaultdict(float), "et_ret": defaultdict(float)}
        for r in run(f"""SELECT supplier, FORMAT_DATE('%Y-%m', doc_date) per, CAST(barcode AS STRING) bc,
                   REGEXP_CONTAINS(LOWER(IFNULL(product_name, '')), r'^сырье') raw,
                   EXTRACT(DAY FROM doc_date) > EXTRACT(DAY FROM LAST_DAY(doc_date)) - 3 last3,
                   SUM(amount_purchase) amt
                 FROM {T('incoming_transactions')}
                 WHERE supplier IN UNNEST(@n) AND doc_date BETWEEN @d1 AND @d2 GROUP BY 1,2,3,4,5""", names_in):
            k = (r["supplier"].strip(), r["per"], r["bc"])
            if r["raw"]:
                out["raw"][k] += f2(r["amt"])
            else:
                out["inc"][k] += f2(r["amt"])
                if r["last3"]:
                    out["inc_last3"][k] += f2(r["amt"])
        for r in run(f"""SELECT supplier, FORMAT_DATE('%Y-%m', doc_date) per, CAST(barcode AS STRING) bc,
                   SUM(amount_purchase) amt FROM {T('outgoing_to_supplier_transactions')}
                 WHERE supplier IN UNNEST(@n) AND doc_date BETWEEN @d1 AND @d2 GROUP BY 1,2,3""", names_ret):
            out["ret"][(r["supplier"].strip(), r["per"], r["bc"])] += f2(r["amt"])
        for r in run(f"""SELECT supplier, FORMAT_DATE('%Y-%m', doc_date) per, SUM(amount) amt
                 FROM {T('torgsoft_incoming_ref_2026')}
                 WHERE supplier IN UNNEST(@n) AND doc_date BETWEEN @d1 AND @d2 GROUP BY 1,2""", names_in):
            out["et_inc"][(r["supplier"].strip(), r["per"])] += f2(r["amt"])
        for r in run(f"""SELECT supplier, FORMAT_DATE('%Y-%m', doc_date) per, SUM(amount_purchase) amt
                 FROM {T('torgsoft_outgoing_ref')}
                 WHERE supplier IN UNNEST(@n) AND doc_date BETWEEN @d1 AND @d2 GROUP BY 1,2""", names_ret):
            out["et_ret"][(r["supplier"].strip(), r["per"])] += f2(r["amt"])
        return out

    # ── агрегаты BQ по поставщику-месяцу ──
    def bq_sum(self, kind, sid, per, bcs=None):
        names = self.aliases[sid]["ret" if kind == "ret" else "in"]
        src = self.bq.get(kind, {})
        return sum(v for (n, p, bc), v in src.items() if p == per and n in names and (bcs is None or bc in bcs))

    def bq_barcodes(self, kind, sid, per):
        names = self.aliases[sid]["ret" if kind == "ret" else "in"]
        agg = defaultdict(float)
        for (n, p, bc), v in self.bq.get(kind, {}).items():
            if p == per and n in names:
                agg[bc] += v
        return agg


# ───────────────────────────── контекст пары ─────────────────────────────
def pair_context(D, t):
    per, members = t["per"], t["members"]
    ctx = {"calcs": [D.calc[(s, per)] for s in members if (s, per) in D.calc]}
    dets = [d for c in ctx["calcs"] for d in D.details.get(c["id"], [])]
    for d in dets:
        c = next(c for c in ctx["calcs"] if c["id"] == d["calculation_id"])
        d["_sid"] = c["supplier_id"]
        d["_rule"] = D.rules.get(d.get("retro_rule_id")) or {}
        net = f2(d["amount_net"])
        d["_rate"] = (f2(d["retro_amount"]) / net) if net else f2(d["applied_percent"]) / 100
    ctx["rule_dets"] = [d for d in dets if d.get("retro_rule_id")]
    ctx["extras"] = [d for d in dets if d.get("adjustment_id")]
    ctx["ship"] = [d for d in ctx["rule_dets"]
                   if (d["_rule"].get("retro_base_type") or "shipment_minus_return") == "shipment_minus_return"
                   and (d["_rule"].get("income_source") or "incoming") == "incoming"]
    ctx["pay"] = [d for d in ctx["rule_dets"] if d["_rule"].get("retro_base_type") == "payments"]
    net = sum(f2(d["amount_net"]) for d in ctx["ship"])
    ctx["net_ship"] = net
    ctx["r_eff"] = (sum(f2(d["retro_amount"]) for d in ctx["ship"]) / net) if net else 0.0
    ctx["det_by_id"] = {d["id"]: d for d in dets}
    return ctx


def cand(code, amount, text, strong=True):
    return {"code": code, "amount": round(amount, 2), "text": text, "strong": strong}


# ───────────────────────────── гипотезы ─────────────────────────────
def h_note(D, t, ctx):
    texts = []
    if t.get("excel_cell") and D.comments.get(t["excel_cell"]):
        texts.append((t["excel_cell"], D.comments[t["excel_cell"]]))
    comps = []
    for cell, txt in texts:
        comps += [(cell, c) for c in components(note_body(txt)) if c["is_money"]]
    non = [(cell, c) for cell, c in comps if any(c["label"].lower().startswith(k) for k in NON_RETRO)]
    if not non:
        for s in t["members"]:
            f = D.facts.get((s, t["per"])) or {}
            for e in f.get("additional_payments") or []:
                if f2(e.get("amount")) > 0:
                    non.append(("админке (доп. выплата в факте)",
                                {"amount": f2(e.get("amount")), "label": e.get("label") or e.get("note") or "доп. выплата"}))
    if not non:
        return []
    got = sum(c["amount"] for _, c in non)
    extra = sum(f2(d["retro_amount"]) for d in ctx["extras"])
    parts = ", ".join(f"{c['amount']:,.0f} {c['label']}" for _, c in non)
    txt = f"в примечании {non[0][0]}: {parts} - это не ретро"
    if extra:
        txt += f"; доплатами в расчёт уже внесено {extra:,.0f}"
    return [cand("NOTE", got - extra, txt)] if abs(got - extra) >= TOL_ABS else []


def h_excel(D, t, ctx):
    X, F, C = t.get("excel_amount"), t.get("fact"), t.get("calc")
    if X is None or F is None or C is None or abs(X - F) < TOL_ABS:
        return []
    if abs(X - C) <= max(TOL_ABS, CLOSE_PCT * abs(C)):
        return [cand("EXCEL", F - X, f"в админке внесено {F:,.0f}, в Excel {t['excel_cell']} = {X:,.0f} - Excel сходится "
                     f"с расчётом, правка в админке")]
    return []


def h_shift(D, t, ctx):
    d = t.get("delta")
    out = []
    for k in (-1, 1):
        pn = month_add(t["per"], k)
        rn = D.recon.get((t["members"][0], pn))
        if not rn or rn.get("delta") is None:
            continue
        dn = f2(rn["delta"])
        if abs(d + dn) <= max(TOL_ABS, 0.1 * abs(d)) and abs(dn) >= TOL_ABS:
            out.append(cand("SHIFT", d, f"гасится месяцем {pn} ({money(dn)}): оплата перенесена между месяцами"))
    return out[:1]


def rule_active(r, per):
    a, b = bounds(per)
    return (not r.get("valid_from") or str(r["valid_from"]) <= str(b)) and \
           (not r.get("valid_to") or str(r["valid_to"]) >= str(a))


def h_threshold(D, t, ctx):
    out = []
    for s in t["members"]:
        for r in D.rules_by_sup[s]:
            thr = f2(r.get("min_purchase_threshold"))
            if not thr or not rule_active(r, t["per"]):
                continue
            excl = set(map(str, r.get("excluded_sku_barcodes") or []))
            inc = sum(v for bc, v in D.bq_barcodes("inc", s, t["per"]).items() if bc not in excl)
            rate = f2(r.get("retro_min")) / 100
            det = next((d for d in ctx["rule_dets"] if d.get("retro_rule_id") == r["id"]), None)
            got = f2(det["retro_amount"]) if det else 0.0
            if inc < thr and got == 0 and t["delta"] > 0:
                out.append(cand("THRESHOLD", inc * rate, f"порог {thr:,.0f}: приход сейчас {inc:,.0f} < порога, расчёт 0, "
                                f"а поставщик заплатил как за {rate:.0%}"))
            elif inc >= thr and got > 0 and t["delta"] < 0 and abs(t["delta"] + got) <= max(TOL_ABS, 0.05 * got):
                out.append(cand("THRESHOLD", -got, f"порог {thr:,.0f}: у нас приход {inc:,.0f} >= порога, "
                                f"поставщик, видимо, считает порог не выполненным"))
    return out


def h_vat(D, t, ctx):
    v = sum(f2(d["retro_amount"]) for d in ctx["rule_dets"] if d["_rule"].get("subtract_vat_from_retro"))
    if not v:
        return []
    return [cand("VAT", v * 0.25, f"НДС не вычтен поставщиком: ретро с НДС-вычетом {v:,.0f} × 1/0,8"),
            cand("VAT", v * (1 / 1.2 / 0.8 - 1), f"поставщик выделил НДС ÷1,2 вместо ×0,8 (+4,2% от {v:,.0f})")]


def h_rate_xls(D, t, ctx):
    cell = t.get("excel_cell") or ""
    row = int(re.sub(r"\D", "", cell) or 0)
    rates = parse_rates(D.conditions.get(row))
    ship = [d for d in ctx["ship"] if f2(d["amount_net"])]
    if not rates or not ship:
        return []
    cur = sorted({round(d["_rate"] * 100, 2) for d in ship})
    if len(cur) != 1 or any(abs(x - cur[0]) < 0.05 for x in rates):
        return []
    out = []
    for rx in rates:
        amt = sum(f2(d["amount_net"]) * (rx / 100) * (0.8 if d["_rule"].get("subtract_vat_from_retro") else 1)
                  - f2(d["retro_amount"]) for d in ship)
        out.append(cand("RATE_XLS", amt, f"в Excel «Условия» {D.conditions.get(row)} -> {rx:g}%, в правиле {cur[0]:g}% "
                        f"(база {ctx['net_ship']:,.0f})", strong=len(rates) == 1))
    return out


def det_barcodes(D, per, det_id):
    return {bc for bc, lst in D.sku_by[per].items() if any(x[1] == det_id for x in lst)}


def h_returns(D, t, ctx):
    out, per = [], t["per"]
    for d in ctx["ship"]:
        pol = d["_rule"].get("returns_policy") or ("subtract" if d["_rule"].get("returns_allowed") else "none")
        name = D.sup.get(d["_sid"], "?")
        if pol == "subtract" and f2(d["amount_returned"]):
            out.append(cand("RETURNS", f2(d["amount_returned"]) * d["_rate"],
                            f"{name}: мы вычли возвраты {f2(d['amount_returned']):,.0f}, поставщик мог не вычесть"))
        elif pol != "subtract" and D.bq:
            ret = D.bq_sum("ret", d["_sid"], per, det_barcodes(D, per, d["id"]))
            if ret:
                out.append(cand("RETURNS", -ret * d["_rate"],
                                f"{name}: политика «{pol}», а в BQ возвраты {ret:,.0f} - поставщик их вычел"))
    return out


def h_sku(D, t, ctx):
    """Сверка SKU расчёта с BigQuery: вне правил, исключённые, сырьё, двойной счёт, дрейф данных."""
    if not D.bq or not ctx["ship"]:
        return []
    out, per = [], t["per"]
    for s in t["members"]:
        mine = [d for d in ctx["ship"] if d["_sid"] == s]
        if not mine:
            continue
        r_s = sum(f2(d["retro_amount"]) for d in mine) / (sum(f2(d["amount_net"]) for d in mine) or 1)
        excluded = {str(b) for r in D.rules_by_sup[s] if rule_active(r, per) for b in (r.get("excluded_sku_barcodes") or [])}
        unc = exc = drift = dbl = 0.0
        unc_n = dbl_n = 0
        for bc, amt in D.bq_barcodes("inc", s, per).items():
            owners = D.sku_by[per].get(bc, [])
            my = [x for x in owners if x[0] == s]
            if my:
                calc_amt = max(x[2] for x in my)
                det = ctx["det_by_id"].get(my[0][1])
                rate = det["_rate"] if det else r_s
                drift += (amt - calc_amt) * rate
                if len(my) > 1:
                    dbl_n += 1
                    dbl += sum(x[2] * (ctx["det_by_id"].get(x[1], {}).get("_rate", r_s)) for x in my[1:])
            elif not owners:
                if bc in excluded:
                    exc += amt
                else:
                    unc += amt
                    unc_n += 1
        raw = D.bq_sum("raw", s, per)
        nm = D.sup.get(s, "?")
        if abs(drift) >= TOL_ABS:
            out.append(cand("SKU_DRIFT", drift, f"{nm}: приход по SKU расчёта в BQ сейчас отличается от сохранённого "
                            f"(перезалив после расчёта) -> {money(drift)} ретро"))
        if dbl:
            out.append(cand("SKU_DOUBLE", -dbl, f"{nm}: {dbl_n} SKU посчитаны в двух правилах -> {money(-dbl)}"))
        if unc * r_s >= TOL_ABS:
            out.append(cand("SKU_OUTSIDE", unc * r_s, f"{nm}: приход {unc:,.0f} по {unc_n} SKU вне всех правил "
                            f"(×{r_s:.1%})"))
        if exc * r_s >= TOL_ABS:
            out.append(cand("SKU_EXCLUDED", exc * r_s, f"{nm}: приход {exc:,.0f} по исключённым SKU (excluded_sku_barcodes) - "
                            f"поставщик мог начислить на них"))
        if raw * r_s >= TOL_ABS:
            out.append(cand("SKU_RAW", raw * r_s, f"{nm}: приход «сырьё» {raw:,.0f} исключён из расчёта"))
    return out


def h_etalon(D, t, ctx):
    if not D.bq or not ctx["ship"]:
        return []
    out, per = [], t["per"]
    for s in t["members"]:
        mine = [d for d in ctx["ship"] if d["_sid"] == s]
        if not mine:
            continue
        names = D.aliases[s]["in"]
        r_s = sum(f2(d["retro_amount"]) for d in mine) / (sum(f2(d["amount_net"]) for d in mine) or 1)
        ours = sum(v for (n, p, bc), v in list(D.bq["inc"].items()) + list(D.bq["raw"].items()) if p == per and n in names)
        et = sum(v for (n, p), v in D.bq["et_inc"].items() if p == per and n in names)
        covered = sum(f2(d["amount_purchased"]) for d in mine)
        share = min(1.0, covered / ours) if ours else 0.0
        if et and abs(et - ours) * share * r_s >= TOL_ABS:
            out.append(cand("ETALON_INC", (et - ours) * share * r_s,
                            f"{D.sup.get(s)}: эталон Торгсофт приход {et:,.0f}, у нас {ours:,.0f} ({money(et - ours)}) "
                            f"-> {money((et - ours) * share * r_s)} ретро"))
        if any((d["_rule"].get("returns_policy") or "") == "subtract" for d in mine):
            rn = D.aliases[s]["ret"]
            o_ret = sum(v for (n, p, bc), v in D.bq["ret"].items() if p == per and n in rn)
            e_ret = sum(v for (n, p), v in D.bq["et_ret"].items() if p == per and n in rn)
            if abs(e_ret - o_ret) * share * r_s >= TOL_ABS:
                out.append(cand("ETALON_RET", -(e_ret - o_ret) * share * r_s,
                                f"{D.sup.get(s)}: эталон возвратов {e_ret:,.0f}, у нас {o_ret:,.0f}"))
    return out


def h_ref(D, t, ctx):
    """income_source=torgsoft_ref: расчёт взял приход из эталона; поставщик мог считать от наших приходов."""
    if not D.bq:
        return []
    out, per = [], t["per"]
    for d in ctx["rule_dets"]:
        if (d["_rule"].get("income_source") or "") != "torgsoft_ref":
            continue
        s = d["_sid"]
        names = D.aliases[s]["in"]
        ours = sum(v for (n, p, bc), v in D.bq["inc"].items() if p == per and n in names)
        et = sum(v for (n, p), v in D.bq["et_inc"].items() if p == per and n in names)
        if abs(ours - et) * d["_rate"] >= TOL_ABS:
            out.append(cand("REF_BASE", (ours - et) * d["_rate"],
                            f"{D.sup.get(s)}: расчёт от эталона Торгсофт {et:,.0f}, наши приходы {ours:,.0f} "
                            f"({money(ours - et)}) - поставщик мог считать от накладных"))
    return out


def h_payments(D, t, ctx):
    out, per = [], t["per"]
    for d in ctx["pay"]:
        rows = D.payments.get(d["_sid"], [])
        by_label = sum(f2(p["amount"]) for p in rows if p.get("period_label") == per)
        by_date = sum(f2(p["amount"]) for p in rows if str(p.get("payment_date") or "")[:7] == per)
        rate = f2(d["_rule"].get("retro_min")) / 100
        if abs(by_date - by_label) * rate >= TOL_ABS:
            out.append(cand("PAYMENTS", (by_date - by_label) * rate,
                            f"{D.sup.get(d['_sid'])}: оплаты по периоду {by_label:,.0f}, по дате платежа {by_date:,.0f}"
                            ))
        if not by_label:
            out.append(cand("PAYMENTS_MISSING", 0, f"{D.sup.get(d['_sid'])}: в supplier_payments нет выписки за {per} - расчёт 0"))
    return out


def h_rate_nice(D, t, ctx, residual):
    ship = [d for d in ctx["ship"] if f2(d["amount_net"])]
    cur = {round(d["_rate"], 4) for d in ship}
    if len(cur) != 1 or not ctx["net_ship"]:
        return []
    r = cur.pop()
    implied = r + residual / ctx["net_ship"]
    nice = round(implied * 200) / 200
    if abs(implied - nice) <= 0.0005 and abs(nice - r) >= 0.004:
        return [cand("RATE_NICE", ctx["net_ship"] * (nice - r), f"возможно: подразумеваемая ставка по факту {implied:.2%} ≈ {nice:.1%} "
                     f"вместо {r:.1%}", strong=False)]
    return []


def h_boundary(D, t, ctx, residual):
    if not D.bq or not ctx["ship"]:
        return []
    per, out = t["per"], []
    for s in t["members"]:
        covered = {bc for bc, lst in D.sku_by[per].items() if any(x[0] == s for x in lst)}
        if residual < 0:
            v = D.bq_sum("inc_last3", s, per, covered) * ctx["r_eff"]
            out.append(cand("BOUNDARY", -v, f"возможно: {D.sup.get(s)}: документы последних 3 дней {per} -> {money(-v)} "
                            f"(поставщик отнёс их к следующему месяцу)", strong=False))
        else:
            pp = month_add(per, -1)
            v = D.bq_sum("inc_last3", s, pp, covered) * ctx["r_eff"]
            out.append(cand("BOUNDARY", v, f"возможно: {D.sup.get(s)}: документы последних 3 дней {pp} -> {money(v)} "
                            f"(поставщик отнёс их к {per})", strong=False))
    return out


# ───────────────────────────── отсутствие данных ─────────────────────────────
def missing_data(D, t, ctx):
    per, st, notes = t["per"], t["status"], []
    names = " + ".join(D.sup.get(s, "?") for s in t["members"])
    inc = sum(D.bq_sum("inc", s, per) for s in t["members"]) if D.bq else None
    et = sum(v for s in t["members"] for (n, p), v in D.bq.get("et_inc", {}).items()
             if p == per and n in D.aliases[s]["in"]) if D.bq else None
    if st == "NOT_PAID":
        paid = sorted({f["period_label"] for s in t["members"] for f in D.facts_all[s] if f2(f["amount_paid"]) > 0})
        before = [p for p in paid if p < per]
        after = [p for p in paid if p > per]
        notes.append(f"начислено {t['calc']:,.0f}; в Excel и в админке за {per} пусто")
        if after:
            notes.append(f"следующие месяцы оплачены ({', '.join(after)}) - {per} пропущен: запросить у поставщика")
        else:
            notes.append(f"после {per} оплат пока нет (последняя оплаченная: {before[-1] if before else '-'}) - "
                         f"проверить, платит ли поставщик вообще")
    elif st == "FACT_NO_CALC":
        act = [r for s in t["members"] for r in D.rules_by_sup[s] if rule_active(r, per)]
        al = sum(len(D.aliases[s]["in"]) for s in t["members"])
        notes.append(f"оплачено {t['fact']:,.0f}, расчёта нет: активных правил {len(act)}, алиасов {al}")
        if inc is not None:
            notes.append(f"приход в BQ по алиасам {inc:,.0f}, в эталоне {et:,.0f}")
            if not inc and et:
                notes.append("в эталоне приход есть, у нас нет - недогруз incoming_transactions")
            elif not inc and not et:
                notes.append("прихода нет нигде - поставщик ведётся под другим именем (алиас решает человек) или платит не за приход")
        if not act:
            notes.append(f"нет активного правила на {per} - завести/продлить правило")
    elif st in ("ZERO_BUT_CALC", "EXCEL_ONLY"):
        if inc is not None:
            notes.append(f"приход в BQ {inc:,.0f}, эталон {et:,.0f}")
        if st == "EXCEL_ONLY":
            notes.append("внести факт в админку (для группы - по каждому поставщику или суммой на основной)")
    return notes


# ───────────────────────────── лестница ─────────────────────────────
LADDER = [h_note, h_excel, h_shift, h_threshold, h_vat, h_rate_xls, h_returns, h_sku, h_etalon, h_ref, h_payments]


def diagnose(D, t):
    ctx = pair_context(D, t)
    d, C = t.get("delta"), abs(t.get("calc") or 0)
    tol = max(TOL_ABS, CLOSE_PCT * C)
    checked, accepted, residual = [], [], d
    if d is not None and abs(d) >= TOL_ABS:
        for h in LADDER:
            cands = h(D, t, ctx)
            checked += cands
            if abs(residual) <= tol:
                continue
            best = None
            for c in cands:
                e = c["amount"]
                if not e or (e > 0) != (residual > 0) or abs(e) > abs(residual) + tol:
                    continue
                if not c["strong"] and abs(residual - e) > tol:
                    continue
                if best is None or abs(residual - e) < abs(residual - best["amount"]):
                    best = c
            if best:
                accepted.append(best)
                residual = round(residual - best["amount"], 2)
        # слабые: круглая ставка - только в «проверено»; граница месяца - только если одна закрывает ВЕСЬ разрыв
        if abs(residual) > tol:
            checked += h_rate_nice(D, t, ctx, residual)
            for c in h_boundary(D, t, ctx, d):
                checked.append(c)
                if not accepted and c["amount"] and (c["amount"] > 0) == (d > 0) and abs(d - c["amount"]) <= tol:
                    accepted.append(c)
                    residual = round(d - c["amount"], 2)
                    break

    lines = missing_data(D, t, ctx)
    if d is not None and abs(d) >= TOL_ABS:
        kind = "ПЕРЕПЛАТА (факт больше расчёта)" if d > 0 else "НЕДОПЛАТА (факт меньше расчёта)"
        closed = abs(residual) <= tol
        if accepted and closed:
            head = f"{kind} {money(d)} объяснена: " + " + ".join(f"[{money(c['amount'])}] {c['text']}" for c in accepted)
            if abs(residual) >= TOL_ABS:
                head += f"; остаток {money(residual)}"
        elif accepted:
            head = (f"{kind} {money(d)}: объяснено {money(d - residual)} - " +
                    " + ".join(f"[{money(c['amount'])}] {c['text']}" for c in accepted) +
                    f"; НЕ объяснено {money(residual)}" + (f" ({residual / C:+.1%})" if C else ""))
        else:
            top = sorted([c for c in checked if c["amount"]], key=lambda c: -abs(c["amount"]))[:3]
            if top:
                tail = "; проверено, не подходит: " + "; ".join(f"[{money(c['amount'])}] {c['text']}" for c in top)
            elif ctx["rule_dets"] and D.bq:
                net_all = sum(f2(x["amount_net"]) for x in ctx["rule_dets"])
                r_all = sum(f2(x["retro_amount"]) for x in ctx["rule_dets"]) / net_all if net_all else 0
                r_impl = (r_all + d / net_all) if net_all else 0
                tail = (f"; данные расчёта совпадают с BigQuery и эталоном (дрейфа, SKU вне правил, дыр нет) - "
                        f"разница в методике поставщика: по факту выходит {r_impl:.2%} от базы {net_all:,.0f} "
                        f"вместо {r_all:.2%}")
            else:
                tail = "; для гипотез нет данных (нет SKU-разбивки расчёта или BigQuery выключен)"
            head = f"{kind} {money(d)}: причина не найдена автоматически" + tail
        lines.insert(0, head)
    else:
        closed = None
    return {"text": sp("; ".join(lines)), "closed": closed, "residual": residual,
            "causes": [{"code": c["code"], "amount": c["amount"], "text": c["text"]} for c in accepted],
            "checked": [{"code": c["code"], "amount": c["amount"]} for c in checked if c["amount"]]}


# ───────────────────────────── запуск ─────────────────────────────
def build_targets(recon_all, supplier=None, period=None, names=None):
    groups = defaultdict(list)
    for r in recon_all:
        if r["status"] not in TARGET or (period and r["period_label"] != period):
            continue
        groups[(r.get("group_key") or r["supplier_id"], r["period_label"])].append(r)
    out = []
    for (_, per), rows in groups.items():
        r0, dd = rows[0], rows[0].get("diagnosis_data") or {}
        members = tuple(sorted(r["supplier_id"] for r in rows))
        if r0.get("group_key"):
            fact, calc = dd.get("group_fact"), dd.get("group_calc")
        else:
            fact = f2(r0["fact_amount"]) if r0.get("fact_amount") is not None else None
            calc = f2(r0["calc_amount"]) if r0.get("calc_amount") is not None else None
        label = " + ".join(names.get(s, "?") for s in members) if names else ""
        if supplier and supplier.lower() not in label.lower():
            continue
        out.append({"members": members, "per": per, "status": r0["status"], "rows": rows, "label": label,
                    "fact": fact, "calc": calc, "delta": f2(r0["delta"]) if r0.get("delta") is not None else None,
                    "excel_amount": f2(r0["excel_amount"]) if r0.get("excel_amount") is not None else None,
                    "excel_cell": r0.get("excel_cell")})
    return sorted(out, key=lambda t: (t["per"], t["label"]))


def run(apply=False, no_bq=False, supplier=None, period=None, verbose=True):
    recon_all = sb.get("retro_reconciliation", "select=*")
    names = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}
    targets = build_targets(recon_all, supplier, period, names)
    if not targets:
        print("[i] Диагностика: проблемных пар нет")
        return []
    D = Data(recon_all, targets, no_bq=no_bq)
    results = []
    for t in targets:
        res = diagnose(D, t)
        results.append((t, res))
    closed = sum(1 for _, r in results if r["closed"])
    with_delta = sum(1 for _, r in results if r["closed"] is not None)
    if verbose:
        print(f"[i] Диагностика: пар {len(targets)} | с расхождением {with_delta}, объяснено полностью {closed} | "
              f"Excel {D.xlsx.name if D.xlsx else '-'} | BQ {'выкл' if no_bq else 'да'}")
        for t, r in results:
            mark = "✔" if r["closed"] else ("…" if r["causes"] else ("✖" if r["closed"] is False else "i"))
            print(f"\n {mark} {t['per']}  {t['label'][:48]:<49} {t['status']:<13} факт {t['fact'] if t['fact'] is not None else '-'}"
                  f"  расчёт {t['calc'] if t['calc'] is not None else '-'}")
            for part in r["text"].split("; "):
                print(f"      {part}")
    dst = OUT / "diagnosis_2026.csv"
    import csv
    with dst.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(["period", "supplier", "status", "fact", "calc", "delta", "closed", "residual", "causes", "diagnosis"])
        for t, r in results:
            w.writerow([t["per"], t["label"], t["status"], t["fact"], t["calc"], t["delta"], r["closed"], r["residual"],
                        " | ".join(f"{c['code']} {c['amount']:+.2f}" for c in r["causes"]), r["text"]])
    print(f"\n[>] {dst.relative_to(OUT.parent)}")

    if apply:
        n = 0
        for t, r in results:
            for row in t["rows"]:
                base = (row.get("diagnosis") or "").split(MARK)[0].rstrip("; ").strip()
                dd = dict(row.get("diagnosis_data") or {})
                dd["diag5"] = {k: r[k] for k in ("closed", "residual", "causes", "checked")}
                sb._req(f"{sb.URL}/rest/v1/retro_reconciliation?id=eq.{row['id']}",
                        {"diagnosis": (f"{base}; " if base else "") + f"{MARK} {r['text']}", "diagnosis_data": dd},
                        "PATCH", {"Prefer": "return=minimal"})
                n += 1
        print(f"[>] retro_reconciliation: диагноз записан в {n} строк")
    return results


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--supplier")
    ap.add_argument("--period")
    ap.add_argument("--no-bq", action="store_true")
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    run(apply=a.apply, no_bq=a.no_bq, supplier=a.supplier, period=a.period)


if __name__ == "__main__":
    main()
