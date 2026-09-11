#!/usr/bin/env python3
"""
import_facts_to_db.py - ЭТАП 3: запись фактов оплат ретро в Supabase.

Принцип: админка - источник истины для уже внесённого, Excel - для нового.
По умолчанию пишутся ТОЛЬКО ячейки, которых нет в retro_payments_fact.
Конфликты (суммы расходятся) показываются, но не перезаписываются - для них нужен
--with-conflicts (или галочка в админке), и прежнее значение сохраняется в amount_previous.

Не пишутся: is_ignored, составные (split_targets) - их месячные ячейки.
Бонусы на открытие магазина пишутся всегда (плательщик = имя строки Excel).

Код разделён: build_plan() - разбор без записи (его же использует сервер админки для черновика),
apply_plan() - запись выбранных строк. CLI печатает план как раньше.

Запуск (из корня):
  python core\\import_facts_to_db.py "Ретро_Excel\\Ретро Бонусы 2026-09-11.xlsx" --sheet 2026
  ... --apply
  ... --with-conflicts --apply
"""
import re, argparse
from datetime import datetime, timezone
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

import sb
from import_retro_facts import norm_name, parse_amount, classify_header, STOP_NAMES
from parse_payment_notes import parse_note, AUTHOR_RE, DATE_RE, EARMARK_RE
from parse_note_components import components, month_refs, ADJ_WORDS
from dump_cell_comments import comments_from_zip

# не ретро в примечании ячейки. «Сгущёнка» сюда НЕ входит: у Юрії это ретро 15% по отдельному правилу (решение 11.09)
NON_RETRO = ("дмп", "кубы", "кубі", "куби", "стойки", "стойка")
RANGE_RE = re.compile(r"\b([а-яіїєґ]{3,})\s*[-—]\s*([а-яіїєґ]{3,})\b", re.IGNORECASE)
SRC = "excel_retro_sheet"
EPS = 1.0  # Excel округлён до гривны: разница < 1 ₴ = совпадение


def note_body(text):
    b = AUTHOR_RE.sub("", re.sub(r"\s+", " ", text or "").strip())
    b = EARMARK_RE.sub(" ", b)
    return DATE_RE.sub(" ", b).strip(" .,-—")


SPREAD_MAX_MONTHS = 6


def find_spread(rows, xl, existing, month_labels, sup):
    """Ручная разноска. Поставщик принёс деньги за 2-3 месяца одной суммой, в Excel она стоит в одном месяце,
    а в админке разнесена по месяцам. По месяцам это «недоплата + переплата», нарастающим итогом - ноль.
    Такой блок - не конфликт: админка права, делать ничего не нужно.
    Блок = подряд идущие месяцы, начинается с месяца, где Excel != админка, и закрывается, когда
    накопленная разница вернулась к нулю (допуск 1 ₴ + копейки админки - Excel округлён до гривны).
    Реальные мелкие разницы (Прометал -68/+65 = -3 ₴) блоком НЕ считаются и остаются конфликтами.
    Пустая ячейка Excel и отсутствие факта в админке = 0."""
    out = []
    for sid in sorted({x[2]["supplier_id"] for x in rows}, key=lambda s: sup.get(s, "")):
        cells = xl.get(sid, {})
        d = []
        for m in month_labels:
            x = cells[m][1] if m in cells else 0.0
            a = existing.get((sid, m))
            d.append((m, x, a, x - (a or 0.0)))
        for i, hit in spread_blocks(d):
            blk = d[i:hit + 1]
            out.append({
                "supplier_id": sid, "supplier": sup.get(sid, "?"),
                "months": [b[0] for b in blk],
                "excel_total": round(sum(b[1] for b in blk), 2),
                "admin_total": round(sum(b[2] or 0.0 for b in blk), 2),
                "cells": [{"period": b[0], "cell": cells[b[0]][0] if b[0] in cells else None,
                           "excel": cells[b[0]][1] if b[0] in cells else None, "admin": b[2]} for b in blk],
                "notes": [note_body(cells[b[0]][2]) for b in blk
                          if b[0] in cells and note_body(cells[b[0]][2])],
            })
    return out


def spread_blocks(d):
    """d = [(месяц, excel, админка|None, excel - админка)] по порядку -> [(i, j)] закрытых блоков ручной разноски."""
    out, i = [], 0
    while i < len(d):
        if abs(d[i][3]) < EPS:
            i += 1
            continue
        cum, tol, hit = 0.0, EPS, None
        for j in range(i, min(i + SPREAD_MAX_MONTHS, len(d))):
            cum += d[j][3]
            tol += abs((d[j][2] or 0.0) - round(d[j][2] or 0.0))  # копейки админки против целых гривен Excel
            signs = {v[3] > 0 for v in d[i:j + 1] if abs(v[3]) >= EPS}
            if len(signs) == 2 and abs(cum) <= tol:
                hit = j  # есть и «недоплата», и «переплата», а в сумме ноль
                break
        if hit is None:
            i += 1
            continue
        out.append((i, hit))
        i = hit + 1
    return out


def build_plan(path, sheet="2026", header_row=1, skip_zero=False):
    """Разбор Excel против текущей БД. Ничего не пишет. -> dict (см. ключи в return)."""
    path = Path(path)
    year = int(re.search(r"(20\d{2})", sheet).group(1))
    ws = load_workbook(path, data_only=True)[sheet]
    now = datetime.now(timezone.utc).isoformat()

    cols = {}
    for c in range(2, ws.max_column + 1):
        cl = classify_header(ws.cell(row=header_row, column=c).value, year)
        if cl and cl[0] != "skip":
            cols[c] = cl

    cmts = {}
    for r in range(header_row + 1, ws.max_row + 1):
        for c in cols:
            cm = ws.cell(row=r, column=c).comment
            if cm and (cm.text or "").strip():
                cmts[(r, c)] = cm.text
    if not cmts:
        for ref, txt in comments_from_zip(path, sheet).items():
            m = re.match(r"([A-Z]+)(\d+)", ref)
            if m:
                cmts[(int(m.group(2)), column_index_from_string(m.group(1)))] = txt

    nmap = {m["excel_name_normalized"]: m for m in sb.get(
        "retro_fact_name_map",
        "select=excel_name_normalized,supplier_id,is_ignored,split_targets,needs_manual")}
    sup = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}
    existing = {(x["supplier_id"], x["period_label"]): float(x["amount_paid"])
                for x in sb.get("retro_payments_fact", "select=supplier_id,period_label,amount_paid")}
    ex_open = {(o["excel_name"], o["store_label"], o["year"]): float(o["amount"])
               for o in sb.get("store_opening_bonuses", "select=excel_name,store_label,year,amount")}

    def opening(sid, raw, label, amt, p, cell):
        prev = ex_open.get((raw, label, year))
        return {"cell": cell, "state": "new" if prev is None else ("same" if abs(prev - amt) < EPS else "changed"),
                "amount_previous": prev,
                "rec": {"supplier_id": sid, "excel_name": raw, "store_label": label, "amount": amt, "year": year,
                        "payment_date": (p["payment_date"].isoformat()
                                         if p["payment_date"] and not p["date_warning"] else None),
                        "source_file": path.name, "imported_at": now}}

    new, conflict, same, openings = [], [], 0, []
    skipped, manual, unmapped, zeros, lost_openings = [], [], set(), 0, []
    totals, sums = {}, defaultdict(float)
    xl = {}  # sid -> {period: (cell, amount, note)} - все месячные ячейки, включая совпадающие
    month_labels = sorted({label for kind, label in cols.values() if kind == "month"})

    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=1).value
        if raw is None or not str(raw).strip():
            continue
        raw = str(raw).strip()
        nm = norm_name(raw)
        if nm in STOP_NAMES:
            for c, (kind, label) in cols.items():
                a = parse_amount(ws.cell(row=r, column=c).value)
                if a is not None:
                    totals[label] = a
            continue
        for c, (kind, label) in cols.items():
            a = parse_amount(ws.cell(row=r, column=c).value)
            if a is not None:
                sums[label] += a

        mp = nmap.get(nm)
        if not mp:
            unmapped.add(raw)
            continue
        if mp.get("is_ignored") or mp.get("split_targets"):
            # Бонус на открытие магазина (колонка = адрес ТТ) платит компания из строки Excel.
            # Плательщик = имя строки (решение 11.09: «Бир Компани (Славутич пиво)»), supplier_id пустой -
            # поставщика не угадываем. Месячные ячейки таких строк не пишутся.
            for c, (kind, label) in cols.items():
                a = parse_amount(ws.cell(row=r, column=c).value)
                if kind == "one_off" and a is not None:
                    cell = f"{get_column_letter(c)}{r}"
                    pn = parse_note(cmts.get((r, c), ""), None, fallback_year=year)
                    openings.append(opening(None, raw, label, a, pn, cell))
                    lost_openings.append((cell, raw, label, a))
            if mp.get("is_ignored"):
                skipped.append(raw)
            else:
                manual.append((raw, "составной платёж"))
            continue
        sid = mp.get("supplier_id")
        if not sid:
            unmapped.add(raw)
            continue

        for c, (kind, label) in cols.items():
            amt = parse_amount(ws.cell(row=r, column=c).value)
            if amt is None:
                continue
            if amt == 0 and skip_zero:
                zeros += 1
                continue
            cell = f"{get_column_letter(c)}{r}"
            txt = cmts.get((r, c), "")
            p = parse_note(txt, label if kind == "month" else None, fallback_year=year)
            body = note_body(txt)

            if kind == "one_off":
                openings.append(opening(sid, raw, label, amt, p, cell))
                continue

            flags = []
            if p["date_warning"]:
                flags.append(p["date_warning"])
            if mp.get("needs_manual"):
                flags.append("поставщик помечен как ручной")

            extra = [{"amount": e["amount"], "label": e["label"], "source": "note"}
                     for e in p["earmarks"]]
            covers, own = [], int(label.split("-")[1])
            if body:
                comps = [x for x in components(body) if x["is_money"]]
                for x in comps:
                    if any(x["label"].startswith(k) for k in NON_RETRO):
                        extra.append({"amount": x["amount"], "label": x["label"], "source": "note"})
                refs = set(month_refs(body))
                rng = RANGE_RE.search(body)
                if rng:
                    a, b = month_refs(rng.group(1)), month_refs(rng.group(2))
                    if a and b:
                        lo, hi = min(a), max(b)
                        if lo < hi:
                            refs |= set(range(lo, hi + 1))
                other = sorted(refs - {own})
                if other and (any(w in body.lower() for w in ADJ_WORDS) or rng):
                    covers = [f"{year}-{m:02d}" for m in other]
                if len(comps) >= 2 and abs(sum(x["amount"] for x in comps) - amt) > 5:
                    flags.append("состав примечания != суммы ячейки")

            rec = {"supplier_id": sid, "period_label": label, "amount_paid": amt,
                   "payment_date": (p["payment_date"].isoformat()
                                    if p["payment_date"] and not p["date_warning"] else None),
                   "notes": txt or None, "source_file": path.name, "import_source": SRC,
                   "imported_at": now, "needs_manual": bool(flags), "amount_previous": None,
                   "covers_periods": covers or None, "additional_payments": extra}

            xl.setdefault(sid, {})[label] = (cell, amt, txt)
            prev = existing.get((sid, label))
            if prev is None:
                new.append((cell, raw, rec, flags))
            elif abs(prev - amt) < EPS:
                same += 1
            else:
                rec["amount_previous"] = prev
                conflict.append((cell, raw, rec, flags))

    spread = find_spread(new + conflict, xl, existing, month_labels, sup)
    drop = {c["cell"] for b in spread for c in b["cells"] if c["cell"]}
    new = [x for x in new if x[0] not in drop]
    conflict = [x for x in conflict if x[0] not in drop]

    control = [{"label": k, "total": v, "rows": round(sums.get(k, 0.0), 2), "ok": abs(sums.get(k, 0.0) - v) < 1}
               for k, v in sorted(totals.items())]
    return {"file": path.name, "sheet": sheet, "year": year, "sup": sup,
            "new": new, "conflict": conflict, "same": same, "openings": openings, "spread": spread,
            "skipped": skipped, "manual": manual, "unmapped": sorted(unmapped), "zeros": zeros,
            "lost_openings": lost_openings, "control": control}


def apply_plan(plan, new_cells=None, conflict_cells=None, openings=True):
    """Запись выбранного. new_cells/conflict_cells - множества адресов ячеек (None = все новые / ни одного конфликта)."""
    batch = [x[2] for x in plan["new"] if new_cells is None or x[0] in new_cells]
    batch += [x[2] for x in plan["conflict"] if conflict_cells and x[0] in conflict_cells]
    written = {"facts": 0, "openings": 0}
    if batch:
        sb.upsert("retro_payments_fact", batch, on_conflict="supplier_id,period_label")
        written["facts"] = len(batch)
    ops = [o["rec"] for o in plan["openings"] if openings and o["state"] != "same"]
    if ops:
        sb.upsert("store_opening_bonuses", ops, on_conflict="excel_name,store_label,year")
        written["openings"] = len(ops)
    return written


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--with-conflicts", action="store_true", help="перезаписать суммы, уже внесённые в админке")
    ap.add_argument("--skip-zero", action="store_true", help="не импортировать нулевые ячейки")
    args = ap.parse_args()

    plan = build_plan(args.xlsx, args.sheet, args.header_row, args.skip_zero)
    sup, new, conflict, openings = plan["sup"], plan["new"], plan["conflict"], plan["openings"]

    bad = [c for c in plan["control"] if not c["ok"]]
    print("[i] Контрольные суммы («Общая сумма»): " + ("все сошлись" if not bad else
          "НЕ СОШЛИСЬ: " + ", ".join(f"{c['label']} итог {c['total']:,.0f} / строки {c['rows']:,.0f}" for c in bad)))
    print(f"[i] Новых к записи: {len(new)} | совпадает: {plan['same']} | конфликтов: {len(conflict)}"
          f" | разнесено вручную (итог сходится): {sum(len(b['cells']) for b in plan['spread'])} мес. в {len(plan['spread'])} блоках")
    print(f"    разовых бонусов: {len(openings)} | ignore: {len(plan['skipped'])} | "
          f"составных: {len(plan['manual'])} | без маппинга: {len(plan['unmapped'])}"
          + (f" | нулей пропущено: {plan['zeros']}" if args.skip_zero else ""))

    if openings:
        byl = defaultdict(float)
        for o in openings:
            byl[o["rec"]["store_label"]] += o["rec"]["amount"]
        print("[i] Бонусы на открытие магазинов (сверить с «Общая сумма»): "
              + ", ".join(f"{k} {v:,.0f}" for k, v in sorted(byl.items()))
              + f" | новых/изменённых: {sum(1 for o in openings if o['state'] != 'same')}")
    if plan["lost_openings"]:
        print(f"[i] из них без поставщика, плательщик = строка Excel ({len(plan['lost_openings'])} шт на "
              f"{sum(x[3] for x in plan['lost_openings']):,.0f}):")
        for cell, raw, label, a in plan["lost_openings"]:
            print(f"    {cell:<6} {raw[:30]:<30} {label:<14} {a:>10,.0f}")

    if plan["spread"]:
        print(f"\n[=] Разнесено вручную по месяцам - итог Excel = итог админки, НЕ конфликты ({len(plan['spread'])}):")
        for b in plan["spread"]:
            print(f"    {b['supplier'][:34]:<34} {b['months'][0]}..{b['months'][-1]}  "
                  f"Excel {b['excel_total']:>11,.0f} = админка {b['admin_total']:>11,.0f}"
                  + (f"  | {' / '.join(b['notes'])[:60]}" if b["notes"] else ""))

    if plan["unmapped"]:
        print("\n[!] Нет в retro_fact_name_map:")
        for x in plan["unmapped"]:
            print(f"    {x}")

    byp = defaultdict(list)
    for cell, raw, rec, fl in new:
        byp[rec["period_label"]].append((cell, raw, rec))
    print("\n[i] Новые записи по периодам:")
    for per in sorted(byp):
        tot = sum(x[2]["amount_paid"] for x in byp[per])
        print(f"    {per}: {len(byp[per]):>2} шт, сумма {tot:>12,.0f}")
        for cell, raw, rec in sorted(byp[per], key=lambda x: sup.get(x[2]["supplier_id"], "")):
            print(f"        {cell:<6} {raw[:30]:<30} -> {sup.get(rec['supplier_id'], '?')[:34]:<34}"
                  f" {rec['amount_paid']:>10,.0f}  {rec['payment_date'] or ''}")

    if conflict:
        print(f"\n[!] Конфликты — НЕ пишутся без --with-conflicts ({len(conflict)}):")
        agg = defaultdict(float)
        for cell, raw, rec, fl in conflict:
            agg[rec["supplier_id"]] += rec["amount_paid"] - rec["amount_previous"]
        for cell, raw, rec, fl in sorted(conflict, key=lambda x: x[2]["period_label"]):
            d = rec["amount_paid"] - rec["amount_previous"]
            tag = "перераспределение?" if abs(agg[rec["supplier_id"]]) < 0.5 else "РАСХОЖДЕНИЕ"
            print(f"    {cell:<6} {raw[:24]:<24} {rec['period_label']}  "
                  f"{rec['amount_previous']:>10,.0f} -> {rec['amount_paid']:>10,.0f} ({d:>+9,.0f})  {tag}")

    flagged = [x for x in new if x[3]]
    if flagged:
        print(f"\n[~] Помечены needs_manual ({len(flagged)}):")
        for cell, raw, rec, fl in flagged:
            print(f"    {cell:<6} {raw[:24]:<24} {rec['period_label']}  {'; '.join(fl)}")

    if not args.apply:
        print("\n[i] Пробный прогон. Для записи добавьте --apply")
        return
    if bad:
        print("\n[!] Контрольные суммы не сошлись - запись отменена")
        return
    w = apply_plan(plan, new_cells=None,
                   conflict_cells={x[0] for x in conflict} if args.with_conflicts else None)
    print(f"[>] retro_payments_fact: записано {w['facts']} | store_opening_bonuses: {w['openings']}")


if __name__ == "__main__":
    main()
