"""Разовый probe: как в админке руками разнесены оплаты за несколько месяцев (Моршин, Галиция, Інтрейд, Рідна марка...).
Excel (сумма + примечание) против админки по месяцам и нарастающим итогом."""
import os, sys
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent; sys.path.insert(0, str(ROOT / "core")); os.chdir(ROOT)
import sb
from import_facts_to_db import build_plan
from openpyxl import load_workbook

XL = ROOT / "Ретро_Excel" / "Ретро Бонусы 2026-09-11.xlsx"
plan = build_plan(XL, "2026")
conf_sids = sorted({x[2]["supplier_id"] for x in plan["conflict"]})
sup = plan["sup"]
facts = {}
for f in sb.get("retro_payments_fact", "select=supplier_id,period_label,amount_paid,notes"):
    facts[(f["supplier_id"], f["period_label"])] = f
ws = load_workbook(XL, data_only=True)["2026"]
# ячейки Excel по поставщику из плана (new/conflict) + совпадающие: пересоберём напрямую
rows = {}
for cell, raw, rec, fl in plan["new"] + plan["conflict"]:
    rows.setdefault(rec["supplier_id"], {})[rec["period_label"]] = (cell, rec["amount_paid"], (rec["notes"] or "").replace("\n", " ")[:70])
import re
from openpyxl.utils import column_index_from_string
for sid in conf_sids:
    cells = rows[sid]
    r = int(re.sub(r"\D", "", next(iter(cells.values()))[0]))
    print(f"\n=== {sup[sid]} (строка {r}): {ws.cell(row=r, column=1).value}")
    cum_x = cum_a = 0.0
    for c in range(3, ws.max_column + 1):
        h = str(ws.cell(row=1, column=c).value or "").strip().lower()
        v = ws.cell(row=r, column=c).value
        per = next((p for p in {x[1] for x in []}), None)
        cm = ws.cell(row=r, column=c).comment
        mon = {"январь": "01", "февраль": "02", "март": "03", "апрель": "04", "май": "05", "июнь": "06",
               "июль": "07", "август": "08", "сентябрь": "09"}.get(h.rstrip("."))
        if not mon:
            continue
        per = f"2026-{mon}"
        a = facts.get((sid, per))
        av = float(a["amount_paid"]) if a else None
        xv = float(v) if isinstance(v, (int, float)) else None
        cum_x += xv or 0
        cum_a += av or 0
        print(f"  {per}  Excel {xv if xv is not None else '-':>10}  админка {av if av is not None else '-':>10}  "
              f"нарастающим: Excel {cum_x:>10,.0f} админка {cum_a:>10,.0f} Δ {cum_x - cum_a:>+9,.0f}  "
              f"| прим: {(cm.text if cm else '').replace(chr(10), ' ')[:60]}")
