#!/usr/bin/env python3
"""
map_excel_to_brands.py - связывает строки Excel с парами (поставщик, бренд)
по УЖЕ существующим справочникам Supabase. Ничего не пишет в БД.

  python map_excel_to_brands.py "Ретро Бонусы.xlsx" --sheet 2026
"""
import re, csv, argparse
from pathlib import Path
from openpyxl import load_workbook

import sb
from import_retro_facts import norm_name, STOP_NAMES

NOISE = ("от оплат", "от оплаты", "от оплати", "кроме", "крім", "ваговий", "весовой")


def split_name(nm):
    brands = " ".join(re.findall(r"\(([^)]*)\)", nm))
    base = re.sub(r"\([^)]*\)", " ", nm)
    tidy = lambda s: re.sub(r"\s+", " ", s).strip(" ,.-")
    return tidy(base), tidy(brands)


def strip_noise(s):
    for n in NOISE:
        s = s.replace(n, " ")
    return re.sub(r"\s+", " ", s).strip(" ,.-")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    args = ap.parse_args()

    ws = load_workbook(Path(args.xlsx), data_only=True)[args.sheet]
    excel = []
    for r in range(args.header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or not str(v).strip():
            continue
        raw = str(v).strip()
        if norm_name(raw) in STOP_NAMES:
            continue
        excel.append((r, raw, norm_name(raw)))

    sups = sb.get("suppliers", "select=id,name")
    brands = sb.get("supplier_brands", "select=id,supplier_id,name")
    aliases = sb.get("supplier_aliases",
                     "select=alias_name,supplier_id,supplier_brand_id,alias_type")
    rules = sb.get("retro_rules",
                   "select=supplier_brand_id,retro_min,retro_base_type,status,valid_from")

    sup_by_id = {s["id"]: s["name"] for s in sups}
    sup_by_norm = {}
    for s in sups:
        sup_by_norm.setdefault(norm_name(s["name"]), s["id"])

    brands_by_sup = {}
    for b in brands:
        brands_by_sup.setdefault(b["supplier_id"], []).append(b)

    rules_by_brand = {}
    for rr in rules:
        if rr.get("status") == "active" or rr.get("status") is None:
            rules_by_brand.setdefault(rr["supplier_brand_id"], []).append(rr)

    # прямой индекс: алиас -> (supplier_id, brand_id)
    alias_idx = {}
    for a in aliases:
        if a.get("alias_type") == "excluded":
            continue
        alias_idx.setdefault(norm_name(a["alias_name"]),
                             (a["supplier_id"], a.get("supplier_brand_id")))

    # индекс "поставщик (бренд)" -> brand_id
    pair_idx = {}
    for b in brands:
        sn = norm_name(sup_by_id.get(b["supplier_id"], ""))
        bn = norm_name(b["name"])
        pair_idx.setdefault(f"{sn} ({bn})", (b["supplier_id"], b["id"]))
        pair_idx.setdefault(bn, (b["supplier_id"], b["id"]))

    print(f"[i] Excel строк: {len(excel)} | поставщиков {len(sups)} | "
          f"брендов {len(brands)} | алиасов {len(aliases)} | правил {len(rules)}")

    rows, stat = [], {}
    for r, raw, nm in excel:
        sid = bid = None
        how = ""
        if nm in alias_idx:
            sid, bid = alias_idx[nm]
            how = "ALIAS"
        elif nm in pair_idx:
            sid, bid = pair_idx[nm]
            how = "PAIR"
        else:
            base, br = split_name(nm)
            base, br = strip_noise(base), strip_noise(br)
            sid = sup_by_norm.get(base)
            if sid:
                cand = brands_by_sup.get(sid, [])
                if br:
                    hit = [b for b in cand if norm_name(b["name"]) == br]
                    if not hit:
                        hit = [b for b in cand
                               if br in norm_name(b["name"]) or norm_name(b["name"]) in br]
                    if len(hit) == 1:
                        bid, how = hit[0]["id"], "SPLIT"
                    elif len(hit) > 1:
                        how = "BRAND_AMBIGUOUS"
                    else:
                        how = "BRAND_NOT_FOUND"
                elif len(cand) == 1:
                    bid, how = cand[0]["id"], "SINGLE_BRAND"
                elif len(cand) > 1:
                    how = "NEED_BRAND"
                else:
                    how = "NO_BRANDS"
            else:
                how = "NO_SUPPLIER"

        n_rules = len(rules_by_brand.get(bid, [])) if bid else 0
        if bid and n_rules == 0:
            how += "+NO_RULE"
        stat[how] = stat.get(how, 0) + 1
        rows.append({
            "row": r, "excel_name": raw, "match": how,
            "supplier": sup_by_id.get(sid, "") if sid else "",
            "supplier_id": sid or "",
            "brand": next((b["name"] for b in brands if b["id"] == bid), "") if bid else "",
            "brand_id": bid or "", "rules": n_rules,
        })

    print("\n[i] Итог:")
    for k in sorted(stat, key=lambda x: -stat[x]):
        print(f"    {k:<20} {stat[k]}")

    bad = [x for x in rows if not x["brand_id"] or "NO_RULE" in x["match"]]
    if bad:
        print(f"\n[!] Требуют внимания ({len(bad)}):")
        for x in bad:
            print(f"    стр {x['row']:>3}  {x['excel_name'][:44]:<46}"
                  f"{x['match']:<20}{x['supplier'][:24]}")

    out = Path(f"excel_brand_map_{args.sheet}.csv")
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(rows)
    print(f"\n[>] {out}")


if __name__ == "__main__":
    main()
