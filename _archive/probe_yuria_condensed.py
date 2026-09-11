"""Разовый probe: Юрія, сгущёнка. Правила из БД, строки расчёта, приход/возвраты сгущёнки в BQ по месяцам,
суммы «сгущёнка» из примечаний Excel -> какой % от базы платит поставщик. -> output/probe_yuria_condensed.txt"""
import os, sys, re
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent; sys.path.insert(0, str(ROOT / "core")); os.chdir(ROOT)
import sb
from google.cloud import bigquery
from openpyxl import load_workbook
from reconcile_facts import newest_excel
from import_facts_to_db import note_body
from parse_note_components import components
from dump_cell_comments import comments_from_zip
from import_retro_facts import classify_header

PROJ = os.environ.get("BQ_PROJECT", "family-market-analytics"); DS = os.environ.get("BQ_DATASET", "family_market")
out = []
sid = [s for s in sb.get("suppliers", "select=id,name") if s["name"] == "Юрія"][0]["id"]
brands = {b["id"]: b for b in sb.get("supplier_brands", f"supplier_id=eq.{sid}&select=id,name")}
rules = sb.get("retro_rules", f"supplier_brand_id=in.({','.join(brands)})&select=*")
out.append("ПРАВИЛА (БД сейчас):")
cond = set()
for r in sorted(rules, key=lambda r: (r["status"], str(r["valid_from"]))):
    out.append(f"  {r['id'][:8]} {brands[r['supplier_brand_id']]['name']} | {r['status']} | {r['valid_from']}..{r.get('valid_to')} | "
               f"{r['retro_min']}% | returns={r.get('returns_policy')} | vat={r.get('subtract_vat_from_retro')} | "
               f"sku={r.get('sku_barcodes')} | excl={r.get('excluded_sku_barcodes')} | {r.get('notes')}")
    for b in (r.get("sku_barcodes") or []) + (r.get("excluded_sku_barcodes") or []):
        cond.add(str(b))
aliases = sorted({a["alias_name"].strip() for a in sb.get("supplier_aliases", f"supplier_id=eq.{sid}&alias_type=neq.excluded&select=alias_name")})
out.append(f"алиасы: {aliases}")

cli = bigquery.Client(project=PROJ)
cfg = bigquery.QueryJobConfig(query_parameters=[bigquery.ArrayQueryParameter("n", "STRING", aliases)])
T = lambda t: f"`{PROJ}.{DS}.{t}`"
q = f"""SELECT CAST(barcode AS STRING) bc, ANY_VALUE(product_name) nm, COUNT(*) n FROM {T('incoming_transactions')}
WHERE TRIM(supplier) IN UNNEST(@n) AND doc_date >= '2026-01-01'
  AND (REGEXP_CONTAINS(LOWER(product_name), r'згущ|сгущ') OR CAST(barcode AS STRING) IN UNNEST({sorted(cond) or ['-']}))
GROUP BY 1 ORDER BY 1"""
skus = {r["bc"]: r["nm"] for r in cli.query(q, job_config=cfg)}
out.append("\nSKU сгущёнки в приходах: " + "; ".join(f"{k} {v}" for k, v in skus.items()))
bcs = sorted(skus) or ["-"]
inc, ret = {}, {}
for tbl, dst in (("incoming_transactions", inc), ("outgoing_to_supplier_transactions", ret)):
    for r in cli.query(f"""SELECT FORMAT_DATE('%Y-%m', doc_date) per, SUM(amount_purchase) a FROM {T(tbl)}
        WHERE TRIM(supplier) IN UNNEST(@n) AND CAST(barcode AS STRING) IN UNNEST({bcs}) AND doc_date >= '2026-01-01'
        GROUP BY 1""", job_config=cfg):
        dst[r["per"]] = float(r["a"] or 0)

calcs = {c["period_label"]: c for c in sb.get("retro_calculations", f"supplier_id=eq.{sid}&select=id,period_label,total_retro,status")}
facts = {f["period_label"]: f for f in sb.get("retro_payments_fact", f"supplier_id=eq.{sid}&select=period_label,amount_paid,notes,additional_payments")}
rid15 = {r["id"] for r in rules if float(r["retro_min"] or 0) == 15.0}
det15 = {}
for per, c in calcs.items():
    for d in sb.get("retro_calculation_details", f"calculation_id=eq.{c['id']}&select=retro_rule_id,amount_purchased,amount_returned,amount_net,retro_amount"):
        if d["retro_rule_id"] in rid15:
            det15[per] = d

xl = newest_excel()
ws = load_workbook(xl, data_only=True)["2026"]
from import_retro_facts import norm_name
mine = {m["excel_name_normalized"] for m in sb.get("retro_fact_name_map", f"supplier_id=eq.{sid}&select=excel_name_normalized")}
row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value and norm_name(str(ws.cell(row=r, column=1).value).strip()) in mine)
cols = {c: classify_header(ws.cell(row=1, column=c).value, 2026) for c in range(2, ws.max_column + 1)}
cm = comments_from_zip(xl, "2026")
out.append(f"\nExcel строка {row}: {ws.cell(row=row, column=1).value} | условия: {ws.cell(row=row, column=2).value}")
out.append(f"\n{'месяц':<8} {'приход сгущ':>11} {'возвр':>8} {'база':>9} | {'расчёт 15%':>10} | {'в примечании':>12} {'% от прихода':>12} {'% от базы':>9} | факт / расчёт всего | примечание")
for c, cl in cols.items():
    if not cl or cl[0] != "month":
        continue
    per = cl[1]
    ref = f"{ws.cell(row=row, column=c).column_letter}{row}"
    txt = note_body(cm.get(ref, "") or (ws.cell(row=row, column=c).comment.text if ws.cell(row=row, column=c).comment else ""))
    note_amt = sum(x["amount"] for x in components(txt) if x["is_money"] and re.search(r"згущ|сгущ", x["label"].lower())) if txt else 0
    i, rt = inc.get(per, 0.0), ret.get(per, 0.0)
    base = i - rt
    d = det15.get(per)
    f, cc = facts.get(per), calcs.get(per)
    out.append(f"{per:<8} {i:>11,.0f} {rt:>8,.0f} {base:>9,.0f} | {(float(d['retro_amount']) if d else 0):>10,.0f} | {note_amt:>12,.0f} "
               f"{(note_amt / i * 100 if i and note_amt else 0):>11.2f}% {(note_amt / base * 100 if base and note_amt else 0):>8.2f}% | "
               f"{f['amount_paid'] if f else '-'} / {cc['total_retro'] if cc else '-'} ({cc['status'] if cc else ''}) | {txt[:70]}")
(ROOT / "output" / "probe_yuria_condensed.txt").write_text("\n".join(out), encoding="utf-8")
print("ok")
