"""Разовый probe (11.09): (1) Маршалл Табако - сколько ТТ проходит порог 8 и 10 SKU в июле/августе;
(2) по каждой паре поставщик-месяц: факт Excel vs сохранённый расчёт vs пересчёт на текущих данных.
Решает по числам, какие месяцы пересчитывать (сохранённый ближе к факту -> НЕ трогать)."""
import os, sys, re, csv
from pathlib import Path
from collections import defaultdict
ROOT = Path(__file__).resolve().parent.parent; sys.path.insert(0, str(ROOT / "core")); os.chdir(ROOT)
import sb
from openpyxl import load_workbook
from import_retro_facts import norm_name, parse_amount, classify_header, STOP_NAMES
from google.cloud import bigquery

P, D = os.environ.get("BQ_PROJECT", "family-market-analytics"), os.environ.get("BQ_DATASET", "family_market")
bq = bigquery.Client(project=P)
XLSX = ROOT / "Ретро_Excel" / "Ретро Бонусы 2026-09-11.xlsx"

# ── (1) Маршалл Табако ──
print("=" * 90 + "\n[1] МАРШАЛЛ ТАБАКО: ТТ по числу SKU поставщика (склады/просрок исключены)")
cols = [r["column_name"] for r in bq.query(f"""SELECT column_name FROM `{P}.{D}.INFORMATION_SCHEMA.COLUMNS`
        WHERE table_name='turnover_monthly' ORDER BY ordinal_position""")]
print("    колонки turnover_monthly:", ", ".join(cols))
sales_col = next((c for c in ("sales_qty", "sold_qty", "qty_sold", "sale_qty") if c in cols), None)
for m in (7, 8):
    defs = {"остаток на начало ИЛИ конец (как в расчёте)": "(t.end_qty > 0 OR t.start_qty > 0)",
            "остаток на конец месяца": "t.end_qty > 0"}
    if sales_col:
        defs[f"были продажи ({sales_col} > 0)"] = f"t.{sales_col} > 0"
    for cap, cond in defs.items():
        rows = list(bq.query(f"""
          WITH sb AS (SELECT DISTINCT CAST(barcode AS STRING) barcode FROM `{P}.{D}.incoming_transactions`
                      WHERE supplier IN ('Маршалл Табако')),
          ss AS (SELECT t.store, COUNT(DISTINCT t.barcode) n FROM `{P}.{D}.turnover_monthly` t
                 JOIN sb ON CAST(t.barcode AS STRING) = sb.barcode
                 WHERE t.year = 2026 AND t.month = {m} AND {cond}
                   AND t.store NOT LIKE '%Склад%' AND t.store NOT LIKE '%Просрок%' GROUP BY t.store)
          SELECT n, COUNT(*) stores FROM ss GROUP BY n ORDER BY n DESC"""))
        dist = {r["n"]: r["stores"] for r in rows}
        ge10 = sum(v for k, v in dist.items() if k >= 10); ge8 = sum(v for k, v in dist.items() if k >= 8)
        print(f"    2026-{m:02d} {cap:<46} >=10 SKU: {ge10:>2} ТТ = {ge10*800:>6,} ₴ | >=8: {ge8:>2} ТТ = {ge8*800:>6,} ₴"
              f" | распределение {dict(sorted(dist.items(), reverse=True))}")
    nsku = list(bq.query(f"""SELECT COUNT(DISTINCT barcode) n FROM `{P}.{D}.incoming_transactions` WHERE supplier='Маршалл Табако'"""))[0]["n"]
    if m == 8:
        print(f"    всего SKU Маршалл в приходах: {nsku}. Факт Excel август: 10 400 = 13 ТТ × 800")

# ── (2) факт vs сохранённое vs пересчёт ──
sup = {s["id"]: s["name"] for s in sb.get("suppliers", "select=id,name")}
byname = {v: k for k, v in sup.items()}
nmap = {m["excel_name_normalized"]: m for m in sb.get("retro_fact_name_map",
        "select=excel_name_normalized,supplier_id,is_ignored,split_targets,needs_manual")}
ws = load_workbook(XLSX, data_only=True)["2026"]
mcols = {}
for c in range(2, ws.max_column + 1):
    cl = classify_header(ws.cell(row=1, column=c).value, 2026)
    if cl and cl[0] == "month":
        mcols[c] = cl[1]
fact, group = {}, {}
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v is None or not str(v).strip() or norm_name(str(v)) in STOP_NAMES:
        continue
    mp = nmap.get(norm_name(str(v).strip()))
    if not mp or mp.get("is_ignored"):
        continue
    if mp.get("split_targets"):
        st = mp["split_targets"]
        ids = []
        for x in (st if isinstance(st, list) else [st]):
            if isinstance(x, dict):
                ids.append(x.get("supplier_id") or byname.get(x.get("name") or x.get("supplier")))
            else:
                ids.append(byname.get(str(x), str(x)))
        key = tuple(sorted(i for i in ids if i))
    else:
        key = (mp["supplier_id"],)
    for c, per in mcols.items():
        a = parse_amount(ws.cell(row=r, column=c).value)
        if a is not None:
            fact[(key, per)] = a

saved, status = defaultdict(float), {}
for c in sb.get("retro_calculations", "select=supplier_id,period_label,total_retro,status"):
    saved[(c["supplier_id"], c["period_label"])] += float(c["total_retro"] or 0)
    status[(c["supplier_id"], c["period_label"])] = c["status"]
extra = defaultdict(float)
cid = {c["id"]: (c["supplier_id"], c["period_label"]) for c in sb.get("retro_calculations", "select=id,supplier_id,period_label")}
for d in sb.get("retro_calculation_details", "retro_rule_id=is.null&select=calculation_id,retro_amount,notes"):
    if "фиксированный ежемесячный бонус" not in (d.get("notes") or "").lower() and d["calculation_id"] in cid:
        extra[cid[d["calculation_id"]]] += float(d["retro_amount"] or 0)
recalc = defaultdict(float)
rows = list(csv.reader(open(ROOT / "_archive" / "recalc_now_2026-01_2026-08.csv", encoding="utf-8-sig"), delimiter=";"))
head = rows[0]
for r in rows[1:]:
    if r and r[0] in byname:
        for p in [h for h in head if re.fullmatch(r"2026-\d\d", h)]:
            v = r[head.index(p)].strip()
            if v:
                recalc[(byname[r[0]], p)] += float(v)

TOL = 1.0
out, verdict = [], defaultdict(lambda: defaultdict(int))
for (key, per), f in sorted(fact.items(), key=lambda x: (x[0][1], sup.get(x[0][0][0], ""))):
    s = sum(saved.get((i, per), 0.0) for i in key)
    n = sum(recalc.get((i, per), 0.0) + extra.get((i, per), 0.0) for i in key)
    ds, dn = f - s, f - n
    st = ",".join(sorted({status.get((i, per), "нет") for i in key}))
    if abs(ds) < TOL and abs(dn) < TOL: v = "оба=факт"
    elif abs(ds) < TOL: v = "СОХР=факт, пересчёт уводит"
    elif abs(dn) < TOL: v = "ПЕРЕСЧЁТ=факт"
    elif abs(ds) <= abs(dn): v = "сохр ближе"
    else: v = "пересчёт ближе"
    verdict[per][v] += 1
    nm = " + ".join(sup.get(i, i) for i in key)
    out.append([nm, per, f, s, n, ds, dn, st, v])

print("\n" + "=" * 90 + "\n[2] ФАКТ Excel vs СОХРАНЁННЫЙ расчёт vs ПЕРЕСЧЁТ сейчас (оба с доп. выплатами), допуск 1 ₴")
for per in sorted(verdict):
    print(f"    {per}: " + " | ".join(f"{k} {v}" for k, v in sorted(verdict[per].items())))
print("\n    Юрія и пары, где пересчёт УВОДИТ от совпавшего факта или ПРИБЛИЖАЕТ к нему (|Δ| >= 100):")
print(f"    {'поставщик':<44}{'период':<9}{'факт':>11}{'сохр':>12}{'пересч':>12}{'факт-сохр':>11}{'факт-пересч':>12}  статус / вердикт")
for nm, per, f, s, n, ds, dn, st, v in out:
    if "Юрія" in nm or (v not in ("оба=факт",) and (abs(ds) >= 100 or abs(dn) >= 100)):
        print(f"    {nm[:43]:<44}{per:<9}{f:>11,.0f}{s:>12,.2f}{n:>12,.2f}{ds:>+11,.0f}{dn:>+12,.0f}  {st} / {v}")
with open(ROOT / "_archive" / "fact_vs_saved_vs_recalc_2026.csv", "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh, delimiter=";")
    w.writerow(["supplier", "period", "fact_excel", "saved", "recalc_now", "fact-saved", "fact-recalc", "status", "verdict"])
    for row in out:
        w.writerow([row[0], row[1]] + [f"{x:.2f}" for x in row[2:7]] + row[7:])
