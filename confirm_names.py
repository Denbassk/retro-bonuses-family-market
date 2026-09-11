#!/usr/bin/env python3
"""
confirm_names.py - ЭТАП 2б: интерактивное подтверждение имён -> retro_fact_name_map.

Клавиши: цифра - выбрать кандидата | s - пропустить | i - не поставщик (ignore)
         /текст - поиск по справочнику | q - сохранить и выйти

Запуск:
  python confirm_names.py "Ретро Бонусы.xlsx" --sheet 2026
"""
import argparse, getpass
from pathlib import Path
from openpyxl import load_workbook

import sb
from import_retro_facts import norm_name, STOP_NAMES
from resolve_supplier_names import score

TOP = 8


def pick(raw, nm, suppliers, sup_norm):
    ranked = sorted(((score(nm, sup_norm[s["id"]]), s) for s in suppliers), key=lambda x: -x[0])
    shown = ranked[:TOP]
    while True:
        print(f"\n{'='*70}\nExcel: {raw}\n  норм: {nm}")
        for i, (sc, s) in enumerate(shown, 1):
            print(f"  {i}. [{sc:.2f}] {s['name']}")
        ans = input("  выбор (цифра / s / i / /поиск / q): ").strip()
        if not ans:
            continue
        if ans == "q":
            return "QUIT"
        if ans == "s":
            return None
        if ans == "i":
            return "IGNORE"
        if ans.startswith("/"):
            q = norm_name(ans[1:])
            found = [s for s in suppliers if q in sup_norm[s["id"]]]
            if not found:
                print("  -- не найдено")
                continue
            shown = [(score(nm, sup_norm[s["id"]]), s) for s in found[:TOP]]
            continue
        if ans.isdigit() and 1 <= int(ans) <= len(shown):
            return shown[int(ans) - 1][1]
        print("  -- не понял")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--sheet", default="2026")
    ap.add_argument("--header-row", type=int, default=1)
    args = ap.parse_args()

    ws = load_workbook(Path(args.xlsx), data_only=True)[args.sheet]
    names = []
    for r in range(args.header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None or not str(v).strip():
            continue
        raw = str(v).strip()
        nm = norm_name(raw)
        if nm in STOP_NAMES or any(n == nm for _, n in names):
            continue
        names.append((raw, nm))

    suppliers = sb.get("suppliers", "select=id,name,legal_name")
    sup_norm = {s["id"]: norm_name(s["name"]) for s in suppliers}
    by_norm = {}
    for s in suppliers:
        by_norm.setdefault(sup_norm[s["id"]], s)

    existing = {m["excel_name_normalized"]
                for m in sb.get("retro_fact_name_map", "select=excel_name_normalized")}
    who = getpass.getuser()

    todo, auto = [], []
    for raw, nm in names:
        if nm in existing:
            continue
        if nm in by_norm:
            auto.append({"excel_name_normalized": nm, "excel_name_raw": raw,
                         "supplier_id": by_norm[nm]["id"], "is_ignored": False,
                         "notes": "EXACT", "confirmed_by": who})
        else:
            todo.append((raw, nm))

    print(f"[i] Уже в карте: {len(existing)} | автоматом EXACT: {len(auto)} | вручную: {len(todo)}")

    new = list(auto)
    for idx, (raw, nm) in enumerate(todo, 1):
        print(f"\n[{idx}/{len(todo)}]", end="")
        res = pick(raw, nm, suppliers, sup_norm)
        if res == "QUIT":
            break
        if res is None:
            continue
        if res == "IGNORE":
            new.append({"excel_name_normalized": nm, "excel_name_raw": raw,
                        "supplier_id": None, "is_ignored": True,
                        "notes": "не поставщик", "confirmed_by": who})
            print("  -> IGNORE")
        else:
            new.append({"excel_name_normalized": nm, "excel_name_raw": raw,
                        "supplier_id": res["id"], "is_ignored": False,
                        "notes": "manual", "confirmed_by": who})
            print(f"  -> {res['name']}")

    if not new:
        print("\n[i] Нечего сохранять.")
        return
    print(f"\n[i] К сохранению: {len(new)}")
    if input("Записать в retro_fact_name_map? (y/n): ").strip().lower() != "y":
        print("[i] Отменено.")
        return
    sb.upsert("retro_fact_name_map", new, on_conflict="excel_name_normalized")
    print(f"[>] Сохранено: {len(new)}")


if __name__ == "__main__":
    main()
